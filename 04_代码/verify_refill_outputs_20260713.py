#!/usr/bin/env python3
"""校验 20260713 补翻输出文件。"""

import json
import os
import re
from pathlib import Path

import openpyxl


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
PROCESS = BASE / "02_过程数据"
FULL_DIR = BASE / "05_上传格式" / "20260713_补翻"
UPLOADABLE_DIR = BASE / "05_上传格式" / "20260713_补翻_可上传"
QC_REPORT = PROCESS / "refill_translation_qc_20260713.json"
VERIFY_REPORT = PROCESS / "refill_output_verify_20260713.json"
MAX_UPLOAD_ROWS = 25000
FORMULA_PREFIXES = ("=", "+", "-", "@")
CYRILLIC_RE = re.compile(r"[\u0400-\u04FF]")


def norm(value):
    return "" if value is None else str(value).strip()


def is_empty(value):
    return norm(value).lower() in {"", "none", "null"}


def verify_file(path, uploadable=False):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    headers = [norm(ws.cell(2, col).value) for col in range(1, 5)]
    lang = headers[2] if len(headers) >= 3 else ""
    data_rows = max(ws.max_row - 2, 0)
    duplicate_keys = 0
    seen = set()
    empty_target_rows = 0
    remark_rows = 0
    formula_cells = 0
    formula_prefix_without_quote = 0
    sr_cyrillic_rows = 0

    for row_idx in range(3, ws.max_row + 1):
        key = norm(ws.cell(row_idx, 1).value)
        target = norm(ws.cell(row_idx, 3).value)
        remark = norm(ws.cell(row_idx, 4).value)
        if key:
            if key in seen:
                duplicate_keys += 1
            seen.add(key)
        if is_empty(target):
            empty_target_rows += 1
        if remark:
            remark_rows += 1
        for col_idx in range(1, 5):
            cell = ws.cell(row_idx, col_idx)
            if cell.data_type == "f":
                formula_cells += 1
            if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES) and not cell.quotePrefix:
                formula_prefix_without_quote += 1
        if lang == "sr" and CYRILLIC_RE.search(target):
            sr_cyrillic_rows += 1

    failures = []
    if headers != ["key ID", "en", lang, "备注"] or not lang:
        failures.append("HEADER_INVALID")
    if ws.max_column != 4:
        failures.append("COLUMN_COUNT_NOT_4")
    if data_rows > MAX_UPLOAD_ROWS:
        failures.append("ROW_LIMIT_EXCEEDED")
    if duplicate_keys:
        failures.append("DUPLICATE_KEY")
    if formula_cells:
        failures.append("FORMULA_CELL")
    if formula_prefix_without_quote:
        failures.append("FORMULA_PREFIX_WITHOUT_QUOTE")
    if uploadable and empty_target_rows:
        failures.append("UPLOADABLE_EMPTY_TARGET")
    if uploadable and remark_rows:
        failures.append("UPLOADABLE_REMARK_NOT_EMPTY")
    if lang == "sr" and sr_cyrillic_rows:
        failures.append("SR_CYRILLIC_REMAINING")

    return {
        "file": str(path),
        "lang": lang,
        "data_rows": data_rows,
        "empty_target_rows": empty_target_rows,
        "remark_rows": remark_rows,
        "duplicate_keys": duplicate_keys,
        "formula_cells": formula_cells,
        "formula_prefix_without_quote": formula_prefix_without_quote,
        "sr_cyrillic_rows": sr_cyrillic_rows,
        "failures": failures,
    }


def main():
    qc = json.loads(QC_REPORT.read_text(encoding="utf-8"))
    expected_full = {
        stats["outputs"][0]: lang
        for lang, stats in qc["langs"].items()
        if stats.get("outputs")
    }
    expected_uploadable = {
        output: lang
        for lang, stats in qc["langs"].items()
        for output in stats.get("uploadable_outputs", [])
    }

    full_files = sorted(FULL_DIR.glob("*.xlsx"))
    uploadable_files = sorted(UPLOADABLE_DIR.glob("*.xlsx"))
    full_results = [verify_file(path, uploadable=False) for path in full_files]
    uploadable_results = [verify_file(path, uploadable=True) for path in uploadable_files]

    all_failures = []
    for scope, results in [("full", full_results), ("uploadable", uploadable_results)]:
        for item in results:
            for failure in item["failures"]:
                all_failures.append(
                    {
                        "scope": scope,
                        "file": item["file"],
                        "lang": item["lang"],
                        "failure": failure,
                    }
                )

    actual_full = {str(path) for path in full_files}
    actual_uploadable = {str(path) for path in uploadable_files}
    missing_full = sorted(set(expected_full) - actual_full)
    extra_full = sorted(actual_full - set(expected_full))
    missing_uploadable = sorted(set(expected_uploadable) - actual_uploadable)
    extra_uploadable = sorted(actual_uploadable - set(expected_uploadable))
    for path in missing_full:
        all_failures.append({"scope": "full", "file": path, "failure": "EXPECTED_FILE_MISSING"})
    for path in extra_full:
        all_failures.append({"scope": "full", "file": path, "failure": "UNEXPECTED_FILE"})
    for path in missing_uploadable:
        all_failures.append({"scope": "uploadable", "file": path, "failure": "EXPECTED_FILE_MISSING"})
    for path in extra_uploadable:
        all_failures.append({"scope": "uploadable", "file": path, "failure": "UNEXPECTED_FILE"})

    report = {
        "full_dir": str(FULL_DIR),
        "uploadable_dir": str(UPLOADABLE_DIR),
        "full_file_count": len(full_files),
        "uploadable_file_count": len(uploadable_files),
        "full_total_rows": sum(item["data_rows"] for item in full_results),
        "uploadable_total_rows": sum(item["data_rows"] for item in uploadable_results),
        "sr_cyrillic_full_rows": sum(item["sr_cyrillic_rows"] for item in full_results),
        "sr_cyrillic_uploadable_rows": sum(item["sr_cyrillic_rows"] for item in uploadable_results),
        "failures": all_failures,
        "full_results": full_results,
        "uploadable_results": uploadable_results,
    }
    VERIFY_REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "verify_report": str(VERIFY_REPORT),
        "full_file_count": report["full_file_count"],
        "uploadable_file_count": report["uploadable_file_count"],
        "full_total_rows": report["full_total_rows"],
        "uploadable_total_rows": report["uploadable_total_rows"],
        "sr_cyrillic_full_rows": report["sr_cyrillic_full_rows"],
        "sr_cyrillic_uploadable_rows": report["sr_cyrillic_uploadable_rows"],
        "failure_count": len(all_failures),
        "failures": all_failures[:20],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
