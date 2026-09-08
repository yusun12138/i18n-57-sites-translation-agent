#!/usr/bin/env python3
"""对 nb 重翻结果做保守占位符/HTML 标签自动修复。

只修复可判断位置的情况：
- 占位符数量一致但名称被翻译：按 EN 顺序替换回原占位符。
- `{0}` 这类数字占位符被翻成独立数字：替换回 `{0}`。
- HTML 标签包裹 EN 中的占位符或数字时：在 NB 译文中同一占位符/数字附近补回标签。
"""

import json
import os
import re
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
CURRENT = BASE / "05_上传格式/20260713_补翻/nb_缺失待平台翻译_20260713.xlsx"
RETRANS_OUTPUT = Path(
    "/Users/10304327/Downloads/nb_%E9%9C%80%E9%87%8D%E7%BF%BB_%E5%B9%B3%E5%8F%B0%E8%BE%93%E5%85%A5_20260714_output.xlsx"
)
RETRANS_DETAIL = BASE / "05_上传格式/20260713_补翻/nb_需重翻_带key明细_20260714.xlsx"

OUT_DIR = BASE / "05_上传格式/20260713_补翻"
OUT_UPLOADABLE = OUT_DIR / "nb_自动修复后_可上传候选_20260714.xlsx"
OUT_STILL_BAD = OUT_DIR / "nb_自动修复后_仍需处理_20260714.xlsx"
OUT_FIX_DETAIL = OUT_DIR / "nb_自动修复明细_20260714.xlsx"
OUT_REPORT = BASE / "02_过程数据/nb_auto_fix_report_20260714.json"

PH_RE = re.compile(r"(\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+)")
TAG_PAIR_RE = re.compile(r"<([a-zA-Z][a-zA-Z0-9]*)([^>]*)>(.*?)</\1>", re.DOTALL)
TAG_RE = re.compile(r"<[^>]+>")
FORMULA_PREFIXES = ("=", "+", "-", "@")

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)


def norm(value):
    return "" if value is None else str(value)


def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", norm(value))


def placeholders(text):
    return PH_RE.findall(norm(text))


def tags(text):
    return TAG_RE.findall(norm(text))


def validate(en, nb):
    issues = []
    if not nb.strip():
        issues.append("EMPTY_NB")
    if sorted(placeholders(en)) != sorted(placeholders(nb)):
        issues.append("PH_MISMATCH")
    en_tags = tags(en)
    nb_tags = tags(nb)
    if en_tags and en_tags != nb_tags:
        issues.append("HTML_TAG_MISMATCH")
    return issues


def set_text(cell, value):
    text = clean(value)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"
    if text.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return list(ws.iter_rows(values_only=True))


def replace_placeholder_sequence(en, nb, actions):
    en_ph = placeholders(en)
    nb_ph = placeholders(nb)
    result = nb
    if en_ph and len(en_ph) == len(nb_ph) and en_ph != nb_ph:
        for old, new in zip(nb_ph, en_ph):
            if old != new:
                result = result.replace(old, new, 1)
        actions.append("PLACEHOLDER_SEQUENCE_RESTORE")
        return result

    if len(en_ph) == 1 and not nb_ph:
        source = en_ph[0]
        m = re.fullmatch(r"\{(\d+)\}", source)
        if m:
            number = re.escape(m.group(1))
            candidate, count = re.subn(rf"(?<!\d){number}(?!\d)", source, result, count=1)
            if count == 1:
                actions.append("NUMERIC_PLACEHOLDER_RESTORE")
                return candidate
    return result


def wrap_first(text, needle_re, start_tag, end_tag):
    match = re.search(needle_re, text)
    if not match:
        return text, False
    start, end = match.span()
    return text[:start] + start_tag + text[start:end] + end_tag + text[end:], True


def restore_simple_html(en, nb, actions):
    if not TAG_PAIR_RE.search(en):
        return nb
    result = nb
    for match in TAG_PAIR_RE.finditer(en):
        tag_name = match.group(1)
        attrs = match.group(2) or ""
        inner = match.group(3) or ""
        start_tag = f"<{tag_name}{attrs}>"
        end_tag = f"</{tag_name}>"
        if start_tag in result and end_tag in result:
            continue

        inner_ph = placeholders(inner)
        wrapped = False
        if inner_ph:
            target = re.escape(inner_ph[0])
            result, wrapped = wrap_first(result, target, start_tag, end_tag)
        else:
            num_match = re.search(r"\d+(?:[.,]\d+)?", inner)
            if num_match:
                target = re.escape(num_match.group(0))
                result, wrapped = wrap_first(result, rf"(?<!\d){target}(?!\d)", start_tag, end_tag)
        if wrapped:
            actions.append(f"HTML_WRAP_{tag_name.upper()}")
    return result


def conservative_fix(en, nb):
    actions = []
    result = nb
    result = replace_placeholder_sequence(en, result, actions)
    result = restore_simple_html(en, result, actions)
    result = replace_placeholder_sequence(en, result, actions)
    return result, actions


def write_upload_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "nb"
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165
    for col, header in enumerate(["key ID", "en", "nb", "备注"], 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 3):
        for col_idx, value in enumerate([row["keyId"], row["en"], row["nb"], ""], 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 28, "B": 70, "C": 70, "D": 20}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def write_detail_workbook(path, rows, title, red=False):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    headers = [
        "key ID", "en", "原nb", "自动修复后nb", "修复动作", "修复前问题",
        "修复后问题", "en占位符", "nb占位符", "en HTML标签", "nb HTML标签", "原文件行号",
    ]
    fill = "F4CCCC" if red else "D9EAD3"
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 2):
        values = [
            row["keyId"], row["en"], row["old_nb"], row["nb"], ";".join(row["actions"]),
            ";".join(row["before_issues"]), ";".join(row["after_issues"]),
            ", ".join(placeholders(row["en"])), ", ".join(placeholders(row["nb"])),
            " | ".join(tags(row["en"])), " | ".join(tags(row["nb"])), row.get("original_row", ""),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {
        "A": 28, "B": 70, "C": 60, "D": 70, "E": 32, "F": 24,
        "G": 28, "H": 30, "I": 30, "J": 44, "K": 44, "L": 12,
    }.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def main():
    current_rows = load_rows(CURRENT)[2:]
    output_rows = load_rows(RETRANS_OUTPUT)[1:]
    detail_rows = load_rows(RETRANS_DETAIL)[1:]
    if len(output_rows) != len(detail_rows):
        raise RuntimeError(f"重翻输出与明细行数不一致: {len(output_rows)} vs {len(detail_rows)}")

    fixed_by_key = {}
    text_mismatch = []
    for output_row, detail_row in zip(output_rows, detail_rows):
        text = norm(output_row[0]).strip()
        new_nb = norm(output_row[1]).strip() if len(output_row) > 1 else ""
        key = norm(detail_row[0]).strip()
        en = norm(detail_row[1]).strip()
        old_nb = norm(detail_row[2]).strip()
        original_row = norm(detail_row[8]).strip() if len(detail_row) > 8 else ""
        if text != en:
            text_mismatch.append({"keyId": key, "text": text[:180], "en": en[:180]})
            continue
        before_issues = validate(en, new_nb)
        candidate, actions = conservative_fix(en, new_nb)
        after_issues = validate(en, candidate)
        fixed_by_key[key] = {
            "keyId": key,
            "en": en,
            "old_nb": old_nb,
            "nb": candidate,
            "actions": actions,
            "before_issues": before_issues,
            "after_issues": after_issues,
            "original_row": original_row,
        }

    uploadable = []
    still_bad = []
    fixed_pass = []
    already_pass = []
    fixed_fail = []
    for current in current_rows:
        key = norm(current[0]).strip()
        en = norm(current[1]).strip()
        nb = norm(current[2]).strip() if len(current) > 2 else ""
        row = fixed_by_key.get(key)
        if row:
            final = row
            if row["after_issues"]:
                still_bad.append(final)
                fixed_fail.append(final)
            else:
                uploadable.append({"keyId": key, "en": en, "nb": row["nb"]})
                fixed_pass.append(final)
        else:
            issues = validate(en, nb)
            synthetic = {
                "keyId": key, "en": en, "old_nb": nb, "nb": nb, "actions": [],
                "before_issues": issues, "after_issues": issues, "original_row": "",
            }
            if issues:
                still_bad.append(synthetic)
            else:
                uploadable.append({"keyId": key, "en": en, "nb": nb})
                already_pass.append(synthetic)

    write_upload_workbook(OUT_UPLOADABLE, uploadable)
    write_detail_workbook(OUT_STILL_BAD, still_bad, "仍需处理", red=True)
    write_detail_workbook(OUT_FIX_DETAIL, fixed_pass + fixed_fail, "自动修复明细", red=False)

    reason_counts = {}
    for row in still_bad:
        key = ";".join(row["after_issues"])
        reason_counts[key] = reason_counts.get(key, 0) + 1

    report = {
        "current_file": str(CURRENT),
        "retranslate_output": str(RETRANS_OUTPUT),
        "retranslate_detail": str(RETRANS_DETAIL),
        "text_mismatch_count": len(text_mismatch),
        "text_mismatch_samples": text_mismatch[:20],
        "current_rows": len(current_rows),
        "retranslate_rows": len(output_rows),
        "already_pass_rows": len(already_pass),
        "fixed_pass_rows": len(fixed_pass),
        "fixed_fail_rows": len(fixed_fail),
        "uploadable_rows": len(uploadable),
        "still_bad_rows": len(still_bad),
        "still_bad_reason_counts": reason_counts,
        "uploadable_file": str(OUT_UPLOADABLE),
        "still_bad_file": str(OUT_STILL_BAD),
        "fix_detail_file": str(OUT_FIX_DETAIL),
    }
    OUT_REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
