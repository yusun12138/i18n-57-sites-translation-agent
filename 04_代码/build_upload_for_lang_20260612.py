#!/usr/bin/env python3
import argparse
import json
import os
import re
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(os.environ.get('I18N_WORKDIR') or Path(__file__).resolve().parents[1]).resolve()
PROCESS = BASE / '02_过程数据'
UPLOAD = BASE / '05_上传格式'
EXCLUDED = BASE / '04_代码' / 'excluded_keyids.txt'
RAW = BASE / '01_原始数据'
DEFAULT_RUN_DATE = os.environ.get('RUN_DATE') or __import__('datetime').datetime.now().strftime('%Y%m%d')


def latest_file(pattern):
    files = sorted(RAW.glob(pattern), key=lambda p: p.stat().st_mtime)
    if not files:
        raise FileNotFoundError(f'未找到源文件: {pattern}')
    return files[-1]


def source_files(run_date):
    return [
        latest_file(f'type1_前端key_platform{platform}_full_{run_date}_*.xlsx')
        for platform in [1, 2, 3, 4, 5]
    ] + [
        latest_file(f'type2_提示语_full_{run_date}_*.xlsx'),
        latest_file(f'type3_错误码_full_{run_date}_*.xlsx'),
    ]

PH_RE = re.compile(r'\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+')
TAG_RE = re.compile(r'<[^>]+>')
ESC_RE = re.compile(r"\\['\u2018\u2019]")
FORMULA_PREFIXES = ('=', '+', '-', '@')
RISK_FILL = PatternFill('solid', fgColor='FFF2CC')
HEADER_FILL = PatternFill('solid', fgColor='D9EAF7')

NOTICE = (
    '模版使用注意事项：\n'
    '1. 模版注意事项内容无需删掉\n'
    '2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n'
    '3. 字段解释：\n'
    'key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n'
    'en——需要更新的key在系统中的原英文文案\n'
    '更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n'
    '4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n'
    '例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n'
    '如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n'
    '5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。'
)


def load_excluded():
    return {
        line.strip()
        for line in EXCLUDED.read_text(encoding='utf-8').splitlines()
        if line.strip() and not line.strip().startswith('#')
    }


def is_empty(value):
    return value is None or str(value).strip() == ''


def norm(value):
    return '' if value is None else str(value).strip()


def qc_issues(en, tr):
    en = str(en or '')
    tr = str(tr or '')
    issues = []
    if not tr.strip():
        issues.append('EMPTY')
    if sorted(PH_RE.findall(en)) != sorted(PH_RE.findall(tr)):
        issues.append('PH_MISMATCH')
    if ESC_RE.findall(en) != ESC_RE.findall(tr):
        issues.append('ESCAPED_SINGLE_QUOTE_MISMATCH')
    if TAG_RE.findall(en) != TAG_RE.findall(tr):
        issues.append('TAG_SEQUENCE_MISMATCH')
    if len(en) >= 1200 and len(tr) < len(en) * 0.45:
        issues.append(f'POSSIBLE_TRUNCATED_LONG_TEXT en_len={len(en)} got_len={len(tr)}')
    return issues


def should_copy_source(en):
    return str(en or '').strip() == "'"


def read_source(lang, files):
    excluded = load_excluded()
    missing_rows = OrderedDict()
    full_cache = {}
    source_stats = []

    for path in files:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = [norm(v).lower() for v in next(rows)]
        key_i = next(i for i, h in enumerate(headers) if h in ('keyid', 'key id', 'key_id'))
        en_i = headers.index('en')
        lang_matches = [i for i, h in enumerate(headers) if h == lang]
        lang_i = lang_matches[-1] if lang == 'id' and lang_matches else (lang_matches[0] if lang_matches else None)
        total = missing = cached = excluded_count = 0

        for row in rows:
            total += 1
            key = norm(row[key_i])
            en = norm(row[en_i])
            tr = '' if lang_i is None else norm(row[lang_i])
            if not key or not en:
                continue
            if key in excluded:
                excluded_count += 1
                continue
            if tr and en not in full_cache and not qc_issues(en, tr):
                full_cache[en] = tr
                cached += 1
            if is_empty(tr) and key not in missing_rows:
                missing_rows[key] = {'key': key, 'en': en, 'source_file': path.name}
                missing += 1
        source_stats.append(
            {
                'file': path.name,
                'rows': total,
                'missing_added': missing,
                'cache_added': cached,
                'excluded_seen': excluded_count,
            }
        )
    return missing_rows, full_cache, source_stats


def load_jsonl(lang, run_date):
    jsonl = PROCESS / f'{lang}_full_{run_date}.jsonl'
    ok = {}
    blocked = {}
    stats = {'ok': 0, 'blocked': 0, 'google': 0, 'dify_fallback': 0, 'lines': 0}
    if not jsonl.exists():
        return ok, blocked, stats

    with open(jsonl, encoding='utf-8') as f:
        for line in f:
            stats['lines'] += 1
            try:
                rec = json.loads(line)
            except Exception:
                continue
            en = rec.get('en_full') or ''
            tr = rec.get('translation') or ''
            status = rec.get('status')
            engine = rec.get('engine')
            if engine in stats:
                stats[engine] += 1
            if status == 'ok' and en and not qc_issues(en, tr):
                ok[en] = tr
                stats['ok'] += 1
            elif en:
                blocked[en] = {'translation': tr, 'issues': rec.get('issues') or ['JSONL_BLOCKED']}
                stats['blocked'] += 1
    return ok, blocked, stats


def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.data_type = 's'
    if cell.value.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def build_workbook(lang, rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = lang
    ws.freeze_panes = 'A3'
    ws.merge_cells('A1:D1')
    set_text(ws['A1'], NOTICE)
    ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 165

    headers = ['key ID', 'en', lang, '备注']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    for row_idx, item in enumerate(rows, 3):
        values = [item['key'], item['en'], item['translation'], item['remark']]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.number_format = '@'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        if item['issues']:
            ws.cell(row_idx, 3).fill = RISK_FILL
            ws.cell(row_idx, 4).fill = RISK_FILL

    widths = {'A': 28, 'B': 62, 'C': 62, 'D': 46}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    UPLOAD.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def verify_upload(path, lang, expected_rows, excluded):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    issues = []
    duplicate_keys = []
    seen = set()
    excluded_keys = []
    empty_target_unmarked = []
    qc_unmarked = []
    formula_cells = []
    formula_prefix_without_quote = []
    marked_rows = 0

    for row_idx in range(3, ws.max_row + 1):
        key = norm(ws.cell(row_idx, 1).value)
        en = '' if ws.cell(row_idx, 2).value is None else str(ws.cell(row_idx, 2).value)
        tr = '' if ws.cell(row_idx, 3).value is None else str(ws.cell(row_idx, 3).value)
        remark = norm(ws.cell(row_idx, 4).value)
        if key in seen:
            duplicate_keys.append(key)
        seen.add(key)
        if key in excluded:
            excluded_keys.append(key)
        if remark:
            marked_rows += 1
        row_qc = qc_issues(en, tr)
        if not tr.strip() and not remark:
            empty_target_unmarked.append(key)
        if row_qc and not remark:
            qc_unmarked.append({'key': key, 'issues': row_qc})
        for col_idx in range(1, 5):
            cell = ws.cell(row_idx, col_idx)
            val = cell.value
            if cell.data_type == 'f':
                formula_cells.append(cell.coordinate)
            if isinstance(val, str) and val.startswith(FORMULA_PREFIXES) and not cell.quotePrefix:
                formula_prefix_without_quote.append(cell.coordinate)

    if ws.max_row - 2 != expected_rows:
        issues.append(f'ROW_COUNT_MISMATCH expected={expected_rows} actual={ws.max_row - 2}')
    if duplicate_keys:
        issues.append(f'DUPLICATE_KEYS count={len(duplicate_keys)} sample={duplicate_keys[:10]}')
    if excluded_keys:
        issues.append(f'EXCLUDED_KEYS_INCLUDED count={len(excluded_keys)} sample={excluded_keys[:10]}')
    if empty_target_unmarked:
        issues.append(f'EMPTY_TARGET_UNMARKED count={len(empty_target_unmarked)} sample={empty_target_unmarked[:10]}')
    if qc_unmarked:
        issues.append(f'QC_UNMARKED count={len(qc_unmarked)} sample={qc_unmarked[:5]}')
    if formula_cells:
        issues.append(f'FORMULA_CELLS count={len(formula_cells)} sample={formula_cells[:10]}')
    if formula_prefix_without_quote:
        issues.append(
            f'FORMULA_PREFIX_WITHOUT_QUOTE count={len(formula_prefix_without_quote)} sample={formula_prefix_without_quote[:10]}'
        )

    return {
        'path': str(path),
        'lang': lang,
        'rows': ws.max_row - 2,
        'marked_rows': marked_rows,
        'duplicate_key_count': len(duplicate_keys),
        'excluded_key_count': len(excluded_keys),
        'empty_target_unmarked_count': len(empty_target_unmarked),
        'qc_unmarked_count': len(qc_unmarked),
        'formula_cell_count': len(formula_cells),
        'formula_prefix_without_quote_count': len(formula_prefix_without_quote),
        'pass': not issues,
        'issues': issues,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--lang', required=True,
                        help='目标语种列名，需要与 Excel 表头保持一致，如 fr、de、lv、el-gr')
    parser.add_argument('--run-date', default=DEFAULT_RUN_DATE,
                        help=f'源文件批次日期，默认取 RUN_DATE 环境变量或当天日期（当前默认 {DEFAULT_RUN_DATE}）')
    args = parser.parse_args()
    lang = args.lang
    out_path = UPLOAD / f'{lang}.xlsx'

    excluded = load_excluded()
    files = source_files(args.run_date)
    missing_rows, full_cache, source_stats = read_source(lang, files)
    jsonl_ok, jsonl_blocked, jsonl_stats = load_jsonl(lang, args.run_date)

    rows = []
    source_counts = {'full_cache': 0, 'jsonl': 0, 'blocked_or_missing': 0}
    for item in missing_rows.values():
        en = item['en']
        translation = ''
        remark = ''
        source = ''
        issues = []

        if should_copy_source(en):
            translation = en
            source = 'jsonl'
        elif en in full_cache:
            translation = full_cache[en]
            source = 'full_cache'
        elif en in jsonl_ok:
            translation = jsonl_ok[en]
            source = 'jsonl'
        elif en in jsonl_blocked:
            translation = jsonl_blocked[en]['translation']
            source = 'blocked_or_missing'
            issues = jsonl_blocked[en]['issues']
        else:
            source = 'blocked_or_missing'
            issues = ['NO_TRANSLATION']

        if not issues:
            issues = qc_issues(en, translation)
        if issues:
            if len(en) >= 5000 or any(issue.startswith('POSSIBLE_TRUNCATED_LONG_TEXT') for issue in issues):
                remark = '长文本疑似未翻译完整，项目最后统一处理；请勿直接上传该行'
            else:
                remark = f"质检风险：{';'.join(issues)}；请勿直接上传该行"

        source_counts[source] += 1
        rows.append(
            {
                'key': item['key'],
                'en': en,
                'translation': translation,
                'remark': remark,
                'issues': issues,
            }
        )

    build_workbook(lang, rows, out_path)
    qc = verify_upload(out_path, lang, len(rows), excluded)
    marked_samples = [
        {
            'key': item['key'],
            'en_len': len(item['en']),
            'translation_len': len(item['translation']),
            'remark': item['remark'],
            'issues': item['issues'],
        }
        for item in rows
        if item['issues']
    ][:50]
    report = {
        'lang': lang,
        'run_date': args.run_date,
        'output': str(out_path),
        'source_files': source_stats,
        'missing_rows': len(missing_rows),
        'full_cache_entries': len(full_cache),
        'jsonl_stats': jsonl_stats,
        'source_counts': source_counts,
        'marked_samples': marked_samples,
        'qc': qc,
    }
    report_path = PROCESS / f'{lang}_upload_qc_{args.run_date}.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
