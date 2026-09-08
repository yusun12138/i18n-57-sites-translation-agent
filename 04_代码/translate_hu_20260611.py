#!/usr/bin/env python3
import hashlib
import json
import os
import random
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
from deep_translator import GoogleTranslator

BASE = Path(os.environ.get('I18N_WORKDIR') or Path(__file__).resolve().parents[1]).resolve()
PROCESS = BASE / '02_过程数据'
JSONL = PROCESS / 'hu_full_20260611.jsonl'
TASK_CACHE = PROCESS / 'hu_full_20260611_tasks.json'
EXCLUDED = BASE / '04_代码' / 'excluded_keyids.txt'

FILES = [
    BASE / '01_原始数据/type1_前端key_platform1_full_20260611_094816.xlsx',
    BASE / '01_原始数据/type1_前端key_platform2_full_20260611_094851.xlsx',
    BASE / '01_原始数据/type1_前端key_platform3_full_20260611_094942.xlsx',
    BASE / '01_原始数据/type1_前端key_platform4_full_20260611_095012.xlsx',
    BASE / '01_原始数据/type1_前端key_platform5_full_20260611_095029.xlsx',
    BASE / '01_原始数据/type2_提示语_full_20260611_095042.xlsx',
    BASE / '01_原始数据/type3_错误码_full_20260611_095054.xlsx',
]

DIFY_PROXY = os.environ.get('DIFY_PROXY') or 'http://127.0.0.1:30001/rest'
DIFY_API_KEY = (os.environ.get('DIFY_API_KEY') or '').strip()
DIFY_HEADERS = {
    'Authorization': f'Bearer {DIFY_API_KEY}',
    'Content-Type': 'application/json',
    'mcp-proxy-path': '/aihelper-dify-api2/v1/workflows/run',
    'mcp-proxy-env': 'prod',
    'mcp-proxy-dc': 'aliyun_dc',
}

PH_RE = re.compile(r'\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+')
TAG_RE = re.compile(r'<[^>]+>')
ESC_RE = re.compile(r"\\['\u2018\u2019]")
CURLY_RE = re.compile(r'\{([^}]*)\}')

GOOGLE_WORKERS = 5
GOOGLE_RETRIES = 2
GOOGLE_INTERVAL = 0.25
FLUSH_EVERY = 500

_google_lock = threading.Lock()
_google_next_at = 0.0
_buffer_lock = threading.Lock()
_pending_records = []


def log(msg):
    print(f'[{datetime.now().strftime("%H:%M:%S")}] {msg}', flush=True)


def md5(text):
    return hashlib.md5(text.encode('utf-8')).hexdigest()


def load_excluded():
    return {
        line.strip()
        for line in EXCLUDED.read_text(encoding='utf-8').splitlines()
        if line.strip() and not line.strip().startswith('#')
    }


def escape_curly(text):
    return CURLY_RE.sub(lambda m: f'[[CURLY_{m.group(1)}]]', text)


def unescape_curly(text):
    return re.sub(r'\[\[CURLY_([^\]]*)\]\]', lambda m: '{' + m.group(1) + '}', text)


def protect(text):
    mapping = []

    def repl(match):
        token = f'ZXPT{len(mapping)}ZX'
        mapping.append((token, match.group(0)))
        return token

    text = TAG_RE.sub(repl, text)
    text = PH_RE.sub(repl, text)
    return text, mapping


def restore(text, mapping):
    for token, raw in mapping:
        text = text.replace(token, raw)
    return text


def qc_issues(en, tr):
    en = str(en or '')
    tr = str(tr or '')
    issues = []
    if not tr.strip():
        issues.append('EMPTY')
    if sorted(PH_RE.findall(en)) != sorted(PH_RE.findall(tr)):
        issues.append('PH_MISMATCH')
    if ESC_RE.findall(en) != ESC_RE.findall(tr):
        issues.append('ESCAPED_SINGLE_QUOTE_MISMATCH')
    if TAG_RE.findall(en) != TAG_RE.findall(tr):
        issues.append('TAG_SEQUENCE_MISMATCH')
    if len(en) >= 1200 and len(tr) < len(en) * 0.45:
        issues.append(f'POSSIBLE_TRUNCATED_LONG_TEXT en_len={len(en)} got_len={len(tr)}')
    return issues


def qc_ok(en, tr):
    return not qc_issues(en, tr)


def wait_google_slot():
    global _google_next_at
    with _google_lock:
        now = time.time()
        if now < _google_next_at:
            time.sleep(_google_next_at - now)
        _google_next_at = time.time() + GOOGLE_INTERVAL + random.uniform(0, 0.06)


def google_translate(en):
    safe, mapping = protect(en)
    last_error = None
    for attempt in range(GOOGLE_RETRIES):
        try:
            wait_google_slot()
            tr = GoogleTranslator(source='en', target='hu').translate(safe)
            tr = restore(str(tr or ''), mapping)
            if tr.strip():
                return tr
            raise RuntimeError('Google returned empty')
        except Exception as exc:
            last_error = exc
            time.sleep(min(8, 1.5 * (attempt + 1) + random.uniform(0.2, 0.8)))
    raise RuntimeError(f'Google failed after {GOOGLE_RETRIES} retries: {last_error}')


def dify_translate(en):
    if not DIFY_API_KEY:
        raise RuntimeError('缺少环境变量 DIFY_API_KEY，不能调用 Dify 翻译接口')
    safe, mapping = protect(en)
    payload = {
        'inputs': {
            'OriginalLanguage': 'English',
            'OriginalText': escape_curly(safe),
            'TargetLanguage': 'hu-HU',
        },
        'response_mode': 'blocking',
        'user': 'hu-google-first-20260611',
    }
    last_error = None
    for attempt in range(2):
        try:
            r = requests.post(
                DIFY_PROXY,
                headers=DIFY_HEADERS,
                data=json.dumps(payload, ensure_ascii=False),
                timeout=120,
            )
            r.raise_for_status()
            out = r.json().get('data', {}).get('outputs', {}).get('output', [])
            for item in out:
                if item.get('target_language_code'):
                    tr = unescape_curly(item.get('translated_text', '') or '')
                    return restore(tr, mapping)
            return ''
        except Exception as exc:
            last_error = exc
            time.sleep(3 * (attempt + 1))
    raise RuntimeError(f'Dify failed: {last_error}')


def flush_records(force=False):
    with _buffer_lock:
        if not _pending_records:
            return 0
        if not force and len(_pending_records) < FLUSH_EVERY:
            return 0
        with open(JSONL, 'a', encoding='utf-8') as f:
            for record in _pending_records:
                f.write(json.dumps(record, ensure_ascii=False) + '\n')
        count = len(_pending_records)
        _pending_records.clear()
        return count


def append_record(record):
    with _buffer_lock:
        _pending_records.append(record)
        should_flush = len(_pending_records) >= FLUSH_EVERY
    if should_flush:
        return flush_records(force=True)
    return 0


def iter_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    ws.reset_dimensions()
    rows = ws.iter_rows(values_only=True)
    headers = [str(v or '').strip() for v in next(rows)]
    lower = [h.lower() for h in headers]
    key_i = next(i for i, h in enumerate(lower) if h in ('keyid', 'key id', 'key_id'))
    en_i = lower.index('en')
    tr_i = lower.index('hu')
    for row in rows:
        key = str(row[key_i] or '').strip()
        if not key:
            continue
        yield key, str(row[en_i] or '').strip(), str(row[tr_i] or '').strip()


def build_needed_and_cache():
    if TASK_CACHE.exists():
        with open(TASK_CACHE, encoding='utf-8') as f:
            data = json.load(f)
        return set(data['needed']), dict(data['cache'])

    excluded = load_excluded()
    cache = {}
    needed = set()
    for path in FILES:
        for key, en, tr in iter_rows(path):
            if key in excluded or not en:
                continue
            if tr and en not in cache and qc_ok(en, tr):
                cache[en] = tr
            if not tr:
                needed.add(en)
    with open(TASK_CACHE, 'w', encoding='utf-8') as f:
        json.dump(
            {'needed': sorted(needed), 'cache': cache},
            f,
            ensure_ascii=False,
        )
    return needed, cache


def load_jsonl_done():
    done = {}
    blocked = {}
    if not JSONL.exists():
        return done, blocked
    with open(JSONL, encoding='utf-8') as f:
        for line in f:
            try:
                rec = json.loads(line)
            except Exception:
                continue
            en = rec.get('en_full') or ''
            tr = rec.get('translation') or ''
            if rec.get('status') == 'ok' and en and qc_ok(en, tr):
                done[en] = tr
            elif rec.get('status') == 'blocked' and en:
                blocked[en] = rec.get('issues') or []
    return done, blocked


def translate_one(en):
    google_error = None
    try:
        tr = google_translate(en)
        issues = qc_issues(en, tr)
        if not issues:
            rec = {
                'en_hash': md5(en),
                'en_full': en,
                'translation': tr,
                'status': 'ok',
                'engine': 'google',
                'ts': datetime.now(timezone.utc).isoformat(),
            }
            append_record(rec)
            return rec
        google_error = ';'.join(issues)
    except Exception as exc:
        google_error = str(exc)

    try:
        tr = dify_translate(en)
        issues = qc_issues(en, tr)
        if not issues:
            rec = {
                'en_hash': md5(en),
                'en_full': en,
                'translation': tr,
                'status': 'ok',
                'engine': 'dify_fallback',
                'google_error': google_error,
                'ts': datetime.now(timezone.utc).isoformat(),
            }
        else:
            rec = {
                'en_hash': md5(en),
                'en_full': en,
                'translation': tr,
                'status': 'blocked',
                'engine': 'dify_fallback',
                'google_error': google_error,
                'issues': issues,
                'ts': datetime.now(timezone.utc).isoformat(),
            }
    except Exception as exc:
        rec = {
            'en_hash': md5(en),
            'en_full': en,
            'translation': '',
            'status': 'blocked',
            'engine': 'dify_fallback',
            'google_error': google_error,
            'issues': [f'DIFY_FAILED: {exc}'],
            'ts': datetime.now(timezone.utc).isoformat(),
        }
    append_record(rec)
    return rec


def main():
    needed, cache = build_needed_and_cache()
    done, blocked = load_jsonl_done()
    # Blocked records are retried; only successful JSONL entries are skipped.
    pending = [en for en in sorted(needed) if en not in cache and en not in done]
    log(f'needed={len(needed)} full_cache={len(cache)} jsonl_done={len(done)} previous_blocked={len(blocked)} pending={len(pending)}')
    ok_count = 0
    blocked_count = 0
    google_count = 0
    fallback_count = 0
    try:
        with ThreadPoolExecutor(max_workers=GOOGLE_WORKERS) as executor:
            futures = {executor.submit(translate_one, en): en for en in pending}
            for i, future in enumerate(as_completed(futures), 1):
                rec = future.result()
                if rec.get('status') == 'ok':
                    ok_count += 1
                    if rec.get('engine') == 'google':
                        google_count += 1
                    else:
                        fallback_count += 1
                else:
                    blocked_count += 1
                if i <= 3 or i % 100 == 0:
                    flushed = flush_records(force=False)
                    suffix = f' flushed={flushed}' if flushed else ''
                    log(f'progress {i}/{len(pending)} ok={ok_count} google={google_count} fallback_ok={fallback_count} blocked={blocked_count}{suffix}')
    finally:
        flushed = flush_records(force=True)
        if flushed:
            log(f'final flush records={flushed}')
    log(f'finished ok={ok_count} google={google_count} fallback_ok={fallback_count} blocked={blocked_count} jsonl={JSONL}')


if __name__ == '__main__':
    main()
