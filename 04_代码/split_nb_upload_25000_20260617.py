#!/usr/bin/env python3
"""把 nb 上传格式文件拆成每份不超过 25000 条数据行。"""

import json
import os
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
INPUT = BASE / "05_上传格式" / "nb_tianshu_byteplus_20260617.xlsx"
OUT_DIR = BASE / "05_上传格式" / "nb_tianshu_byteplus_20260617_split_25000"
MANIFEST = OUT_DIR / "manifest.json"
MAX_DATA_ROWS = 25_000
FORMULA_PREFIXES = ("=", "+", "-", "@")

NOTICE_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RISK_FILL = PatternFill("solid", fgColor="FFF2CC")


def set_text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    if cell.value.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def write_part(rows, part_no):
    wb = Workbook()
    ws = wb.active
    ws.title = "nb"
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")

    set_text(ws["A1"], rows["notice"])
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].fill = NOTICE_FILL
    ws.row_dimensions[1].height = 165

    for col_idx, header in enumerate(rows["headers"], 1):
        cell = ws.cell(2, col_idx)
        set_text(cell, header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    risk_rows = 0
    for out_row_idx, values in enumerate(rows["data"], 3):
        has_risk = bool(values[3])
        if has_risk:
            risk_rows += 1
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(out_row_idx, col_idx)
            set_text(cell, value)
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if has_risk and col_idx in (3, 4):
                cell.fill = RISK_FILL

    for col, width in {"A": 28, "B": 62, "C": 62, "D": 46}.items():
        ws.column_dimensions[col].width = width

    path = OUT_DIR / f"nb_tianshu_byteplus_20260617_part{part_no:02d}.xlsx"
    wb.save(path)
    return path, risk_rows


def verify(path, expected_rows):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb.active
    headers = [ws.cell(row=2, column=i).value for i in range(1, 5)]
    if headers != ["key ID", "en", "nb", "备注"]:
        raise ValueError(f"{path.name} 表头异常: {headers}")
    data_rows = ws.max_row - 2
    if data_rows != expected_rows:
        raise ValueError(f"{path.name} 行数异常: expected={expected_rows}, actual={data_rows}")
    if data_rows > MAX_DATA_ROWS:
        raise ValueError(f"{path.name} 超过 {MAX_DATA_ROWS} 行: {data_rows}")
    return {"path": str(path), "rows": data_rows, "bytes": path.stat().st_size}


def main():
    src_wb = openpyxl.load_workbook(INPUT, read_only=True, data_only=False)
    src_ws = src_wb.active
    notice = src_ws.cell(1, 1).value or ""
    headers = [src_ws.cell(2, col).value for col in range(1, 5)]
    if headers != ["key ID", "en", "nb", "备注"]:
        raise ValueError(f"源文件表头异常: {headers}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    files = []
    chunk = []
    part_no = 1
    total_rows = 0
    total_risk_rows = 0

    for row in src_ws.iter_rows(min_row=3, max_col=4, values_only=True):
        chunk.append(["" if value is None else str(value) for value in row])
        if len(chunk) >= MAX_DATA_ROWS:
            payload = {"notice": notice, "headers": headers, "data": chunk}
            path, risk_rows = write_part(payload, part_no)
            item = verify(path, len(chunk))
            item["risk_rows"] = risk_rows
            files.append(item)
            total_rows += len(chunk)
            total_risk_rows += risk_rows
            part_no += 1
            chunk = []

    if chunk:
        payload = {"notice": notice, "headers": headers, "data": chunk}
        path, risk_rows = write_part(payload, part_no)
        item = verify(path, len(chunk))
        item["risk_rows"] = risk_rows
        files.append(item)
        total_rows += len(chunk)
        total_risk_rows += risk_rows

    expected_rows = src_ws.max_row - 2
    if total_rows != expected_rows:
        raise ValueError(f"拆分总行数异常: expected={expected_rows}, actual={total_rows}")

    manifest = {
        "input": str(INPUT),
        "max_data_rows": MAX_DATA_ROWS,
        "total_data_rows": total_rows,
        "total_risk_rows": total_risk_rows,
        "files": files,
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
