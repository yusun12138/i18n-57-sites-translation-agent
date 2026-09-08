#!/usr/bin/env python3
"""全面排查线上全量 nb 仍像英文的内容。

检测重点：
- en 与 nb 完全相同
- 全角/半角、空格、标点归一后相同
- nb 中残留明显英文 UI 词/短语
- 单独分出主题词、品牌、代码、枚举，避免误报混入高置信清单
"""

import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
DATA_DIR = BASE / "01_原始数据"
REPORT_DIR = BASE / "03_分析报告"
PROCESS_DIR = BASE / "02_过程数据"
RUN_DATE = "20260717"

PH_RE = re.compile(r"(\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+|\$\{[^}]+\})")
TAG_RE = re.compile(r"<[^>]+>")
NON_LATIN_RE = re.compile(r"[\u4e00-\u9fff\u3040-\u30ff\u0400-\u04ff\u0600-\u06ff]")
URL_RE = re.compile(
    r"(https?://|www\.|\.com\b|\.net\b|\.cn\b|\.jpg\b|\.jpeg\b|\.png\b|"
    r"\.webp\b|\.gif\b|\.svg\b|\.mp4\b|/assets?/|^//)",
    re.I,
)

BRAND_WORDS = {
    "shein", "romwe", "motf", "dazy", "sheglam", "evolushein", "shein.com",
    "paypal", "klarna", "afterpay", "tabby", "tamara", "apple", "google",
    "facebook", "instagram", "tiktok", "youtube", "dji", "nintendo", "visa",
    "mastercard", "maestro", "amex", "american", "express",
}

CODE_WORDS = {
    "faq", "pin", "sms", "sku", "id", "uid", "url", "uri", "api", "pc",
    "app", "ios", "android", "vip", "cod", "cvv", "iban", "bic", "gb",
    "mb", "kb", "eur", "usd", "brl", "mxn", "lv", "iof", "cpf", "cnpj",
    "dni", "cuit",
}

MONTH_WORDS = {
    "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "sept",
    "oct", "nov", "dec",
}

THEME_WORDS = {
    "ootd", "chic", "punk", "goth", "emo", "grunge", "buzz", "spice",
    "leather", "western", "folk", "runway", "club", "gals", "haul",
    "live", "era", "rock", "freak", "camis", "parcelado",
}

SHARED_OR_LOAN_WORDS = {
    "design", "status", "album", "video", "filter", "chat", "live", "bio",
    "email", "e-mail", "shopping", "software", "update", "sofort",
}

STRONG_ENGLISH_WORDS = {
    "selected", "select", "refund", "processing", "shipping", "ship", "free",
    "order", "invite", "invites", "get", "instant", "deal", "deals",
    "valued", "user", "congrats", "congratulations", "coupon", "bundle",
    "random", "color", "item", "items", "save", "verification", "method",
    "return", "returns", "address", "password", "sign", "checkout",
    "payment", "delivery", "failed", "successful", "success", "error",
    "apply", "cancel", "continue", "remove", "delete", "submit", "search",
    "view", "details", "wallet", "points", "reward", "rewards", "price",
    "total", "quantity", "size", "help", "message", "notification", "login",
    "register", "download", "upload", "allow", "access", "physical", "gift",
    "cards", "certified", "quality", "global", "authority", "details",
}

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HIGH_FILL = PatternFill("solid", fgColor="F4CCCC")
MID_FILL = PatternFill("solid", fgColor="FFF2CC")
LOW_FILL = PatternFill("solid", fgColor="D9EAD3")


def log(message):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def norm_cell(value):
    return "" if value is None else str(value).strip()


def nfkc(text):
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return text


def normalize_for_equal(text):
    text = nfkc(text).lower()
    text = TAG_RE.sub("", text)
    text = PH_RE.sub(lambda match: "{" + re.sub(r"\s+", "", match.group(0).lower()) + "}", text)
    text = re.sub(r"\s+", "", text)
    text = text.replace("，", ",").replace("。", ".").replace("：", ":")
    return text


def normalize_alpha(text):
    text = nfkc(text).lower()
    text = TAG_RE.sub(" ", text)
    text = PH_RE.sub(" ", text)
    text = re.sub(r"\b(shein|romwe|motf|dazy|sheglam)\b", " ", text)
    return text


def words(text):
    return re.findall(r"[a-zA-Z]+", normalize_alpha(text))


def is_url_or_asset(text):
    return bool(URL_RE.search(text or ""))


def is_code_like(text, word_list):
    raw = nfkc(text).strip()
    if not raw:
        return False
    compact = re.sub(r"[^A-Za-z0-9]", "", raw)
    if len(compact) <= 2:
        return True
    if re.fullmatch(r"[A-Z0-9._/\-]{1,12}", raw):
        return True
    low_words = {word.lower() for word in word_list}
    if low_words and low_words <= CODE_WORDS:
        return True
    if re.search(r"\b(Lv\.?\d+|VIP\d*|\d+\s*(GB|MB|KB))\b", raw, re.I):
        return True
    return False


def classify(en, nb):
    en = norm_cell(en)
    nb = norm_cell(nb)
    if not nb:
        return "空值", "HIGH", "NB_EMPTY"
    if is_url_or_asset(nb):
        return "URL/资源路径", "LOW", "URL_OR_ASSET"
    if NON_LATIN_RE.search(en) and not re.search(r"[A-Za-z]", en):
        return "非英文源文案", "LOW", "NON_ENGLISH_SOURCE"

    en_norm = normalize_for_equal(en)
    nb_norm = normalize_for_equal(nb)
    nb_words = words(nb)
    nb_word_set = {word.lower() for word in nb_words}

    if en_norm and nb_norm and en_norm == nb_norm:
        if is_code_like(nb, nb_words):
            return "代码/枚举", "LOW", "NORM_EQUAL_CODE"
        if nb_word_set and nb_word_set <= BRAND_WORDS:
            return "品牌/支付/专名", "LOW", "NORM_EQUAL_BRAND"
        if nb_word_set and (nb_word_set & THEME_WORDS) and len(nb_word_set) <= 4:
            return "主题词/活动名/专有名", "LOW", "NORM_EQUAL_THEME"
        return "归一化后等于英文", "HIGH", "NORM_EQUAL_EN"

    if is_code_like(nb, nb_words):
        return "代码/枚举", "LOW", "CODE_OR_ENUM"

    if nb_word_set and nb_word_set <= BRAND_WORDS:
        return "品牌/支付/专名", "LOW", "BRAND_OR_PAYMENT"

    if nb_word_set and (nb_word_set & MONTH_WORDS):
        return "月份/单位/等级", "LOW", "MONTH_UNIT_LEVEL"

    if nb_word_set and (nb_word_set & THEME_WORDS) and len(nb_word_set) <= 5:
        return "主题词/活动名/专有名", "LOW", "THEME_OR_PROPER_NOUN"

    strong_hits = sorted(nb_word_set & STRONG_ENGLISH_WORDS)
    if len(strong_hits) >= 2:
        return "明显英文短语", "HIGH", "STRONG_ENGLISH_PHRASE:" + ",".join(strong_hits[:8])
    if len(strong_hits) == 1:
        word = strong_hits[0]
        if word in {"selected", "refund", "shipping", "verification", "password", "checkout", "payment", "delivery"}:
            return "明显英文单词", "HIGH", "STRONG_ENGLISH_WORD:" + word
        return "疑似英文 UI 单词", "MEDIUM", "ENGLISH_UI_WORD:" + word

    if len(nb_word_set) <= 3 and nb_word_set & SHARED_OR_LOAN_WORDS:
        return "可能是挪威语借词/主题词", "LOW", "SHARED_LOAN_WORD"

    return "", "", ""


def latest_export_files(timestamp=None):
    if timestamp:
        files = sorted(DATA_DIR.glob(f"*full_{timestamp}.xlsx"))
        if not files:
            raise FileNotFoundError(f"找不到 timestamp={timestamp} 的全量文件")
        return timestamp, files
    groups = defaultdict(list)
    for path in DATA_DIR.glob("*full_*.xlsx"):
        match = re.search(r"_full_(\d{8}_\d{6})\.xlsx$", path.name)
        if match:
            groups[match.group(1)].append(path)
    if not groups:
        raise FileNotFoundError("没有找到全量导出文件")
    ts = sorted(groups)[-1]
    return ts, sorted(groups[ts])


def load_history_candidates():
    candidates = defaultdict(list)
    history_files = [
        BASE / "05_上传格式/nb.xlsx",
        BASE / "05_上传格式/nb1.xlsx",
        BASE / "05_上传格式/nb2.xlsx",
        BASE / "05_上传格式/20260716_nb_sr_英文残留重翻/nb_英文残留重翻_可上传_20260716.xlsx",
    ]
    for path in history_files:
        if not path.exists():
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                first = [norm_cell(v) for v in next(rows)]
            except StopIteration:
                continue
            header = first
            if "key ID" not in header and "keyId" not in header:
                try:
                    header = [norm_cell(v) for v in next(rows)]
                except StopIteration:
                    continue
            key_col = None
            for h in ("key ID", "keyId", "keyID"):
                if h in header:
                    key_col = header.index(h)
                    break
            if key_col is None or "nb" not in header:
                continue
            nb_col = header.index("nb")
            en_col = header.index("en") if "en" in header else None
            for row in rows:
                key = norm_cell(row[key_col] if key_col < len(row) else "")
                nb = norm_cell(row[nb_col] if nb_col < len(row) else "")
                en = norm_cell(row[en_col] if en_col is not None and en_col < len(row) else "")
                if not key or not nb:
                    continue
                category, confidence, reason = classify(en or nb, nb)
                if confidence != "HIGH" and reason not in {"CODE_OR_ENUM", "NORM_EQUAL_CODE"}:
                    candidates[key].append({"nb": nb, "source": path.name})
    return candidates


def scan(timestamp=None):
    timestamp, files = latest_export_files(timestamp)
    history = load_history_candidates()
    records = []
    summary = Counter()
    confidence_counts = Counter()

    for path in files:
        log(f"扫描 {path.name}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            rows = ws.iter_rows(values_only=True)
            try:
                headers = [norm_cell(v) for v in next(rows)]
            except StopIteration:
                continue
            if "keyId" not in headers or "en" not in headers or "nb" not in headers:
                continue
            key_col = headers.index("keyId")
            en_col = headers.index("en")
            nb_col = headers.index("nb")
            zh_col = headers.index("zh") if "zh" in headers else None
            for row_num, row in enumerate(rows, start=2):
                key = norm_cell(row[key_col] if key_col < len(row) else "")
                en = norm_cell(row[en_col] if en_col < len(row) else "")
                nb = norm_cell(row[nb_col] if nb_col < len(row) else "")
                if not key or not en:
                    continue
                category, confidence, reason = classify(en, nb)
                if not category:
                    continue
                zh = norm_cell(row[zh_col] if zh_col is not None and zh_col < len(row) else "")
                history_nb = ""
                history_src = ""
                if key in history:
                    history_nb = history[key][0]["nb"]
                    history_src = history[key][0]["source"]
                record = {
                    "置信度": confidence,
                    "分类": category,
                    "原因": reason,
                    "key ID": key,
                    "en": en,
                    "当前nb": nb,
                    "zh": zh,
                    "历史可参考nb": history_nb,
                    "历史来源": history_src,
                    "source_file": path.name,
                    "sheet": ws.title,
                    "row": row_num,
                    "en_norm": normalize_for_equal(en),
                    "nb_norm": normalize_for_equal(nb),
                }
                records.append(record)
                summary[(confidence, category, reason.split(":")[0])] += 1
                confidence_counts[confidence] += 1

    output_xlsx = REPORT_DIR / f"nb_线上英文残留全面排查_{RUN_DATE}.xlsx"
    output_json = PROCESS_DIR / f"nb_english_residual_full_scan_{RUN_DATE}.json"
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    PROCESS_DIR.mkdir(parents=True, exist_ok=True)

    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "summary"
    ws.append(["timestamp", timestamp])
    ws.append(["generated_at", datetime.now(timezone.utc).isoformat()])
    ws.append(["source_files", len(files)])
    ws.append([])
    ws.append(["置信度", "分类", "原因", "数量"])
    for (confidence, category, reason), count in sorted(summary.items(), key=lambda item: (item[0][0], item[0][1], item[0][2])):
        ws.append([confidence, category, reason, count])
    ws.append([])
    ws.append(["置信度汇总", "数量"])
    for confidence, count in confidence_counts.most_common():
        ws.append([confidence, count])

    headers = ["置信度", "分类", "原因", "key ID", "en", "当前nb", "zh", "历史可参考nb", "历史来源", "source_file", "sheet", "row"]
    sheet_specs = [
        ("高置信需处理", lambda r: r["置信度"] == "HIGH", HIGH_FILL),
        ("中置信人工确认", lambda r: r["置信度"] == "MEDIUM", MID_FILL),
        ("低置信主题专名代码", lambda r: r["置信度"] == "LOW", LOW_FILL),
        ("全部命中", lambda r: True, HEADER_FILL),
    ]
    for title, predicate, fill in sheet_specs:
        ws = wb_out.create_sheet(title)
        ws.append(headers)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(1, col_idx)
            cell.font = Font(bold=True)
            cell.fill = fill if title != "全部命中" else HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
        for record in records:
            if not predicate(record):
                continue
            ws.append([record.get(h, "") for h in headers])
        for col, width in {
            "A": 12, "B": 24, "C": 34, "D": 30, "E": 70, "F": 70,
            "G": 40, "H": 70, "I": 24, "J": 48, "K": 18, "L": 10,
        }.items():
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb_out.save(output_xlsx)

    report = {
        "timestamp": timestamp,
        "files": [str(path) for path in files],
        "output_xlsx": str(output_xlsx),
        "confidence_counts": dict(confidence_counts),
        "summary": [
            {"confidence": c, "category": cat, "reason": reason, "count": count}
            for (c, cat, reason), count in summary.items()
        ],
        "total_records": len(records),
    }
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main():
    report = scan()
    log(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
