#!/usr/bin/env python3
"""将 sr 上传文件中的塞尔维亚西里尔字母转写为拉丁字母，并生成质检报告。"""

import json
import os
import re
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
UPLOAD = BASE / "05_上传格式"
PROCESS = BASE / "02_过程数据"
BACKUP_DIR = UPLOAD / "backup_sr_cyrillic_20260622"
SOURCE = UPLOAD / "sr.xlsx"
MAX_DATA_ROWS = 25_000
FORMULA_PREFIXES = ("=", "+", "-", "@")
CYRILLIC_RE = re.compile(r"[\u0400-\u04FF]")

SR_CYR_TO_LAT = {
        "А": "A",
        "Б": "B",
        "В": "V",
        "Г": "G",
        "Д": "D",
        "Ђ": "Đ",
        "Е": "E",
        "Ж": "Ž",
        "З": "Z",
        "И": "I",
        "Ј": "J",
        "К": "K",
        "Л": "L",
        "Љ": "Lj",
        "М": "M",
        "Н": "N",
        "Њ": "Nj",
        "О": "O",
        "П": "P",
        "Р": "R",
        "С": "S",
        "Т": "T",
        "Ћ": "Ć",
        "У": "U",
        "Ф": "F",
        "Х": "H",
        "Ц": "C",
        "Ч": "Č",
        "Џ": "Dž",
        "Ш": "Š",
        "а": "a",
        "б": "b",
        "в": "v",
        "г": "g",
        "д": "d",
        "ђ": "đ",
        "е": "e",
        "ж": "ž",
        "з": "z",
        "и": "i",
        "ј": "j",
        "к": "k",
        "л": "l",
        "љ": "lj",
        "м": "m",
        "н": "n",
        "њ": "nj",
        "о": "o",
        "п": "p",
        "р": "r",
        "с": "s",
        "т": "t",
        "ћ": "ć",
        "у": "u",
        "ф": "f",
        "х": "h",
        "ц": "c",
        "ч": "č",
        "џ": "dž",
        "ш": "š",
        # Google/Dify 偶尔混入非塞尔维亚西里尔字符，按最接近拉丁写法处理。
        "Ё": "Jo",
        "Й": "J",
        "Щ": "Šč",
        "Ъ": "",
        "Ы": "Y",
        "Ь": "",
        "Э": "E",
        "Ю": "Ju",
        "Я": "Ja",
        "ё": "jo",
        "й": "j",
        "щ": "šč",
        "ъ": "",
        "ы": "y",
        "ь": "",
        "э": "e",
        "ю": "ju",
        "я": "ja",
}

CYR_UPPER = set("АБВГДЂЕЖЗИЈКЛЉМНЊОПРСТЋУФХЦЧЏШЁЙЩЪЫЬЭЮЯ")
CYR_LOWER = set("абвгдђежзијклљмнњопрстћуфхцчџшёйщъыьэюя")
DIGRAPH_UPPER = {"Љ": "LJ", "Њ": "NJ", "Џ": "DŽ"}


def set_text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    cell.number_format = "@"
    if cell.value.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def copy_cell(src, dst):
    if isinstance(src, MergedCell):
        return
    dst.value = src.value
    dst.data_type = src.data_type
    dst.number_format = src.number_format
    dst.quotePrefix = src.quotePrefix
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    if isinstance(dst.value, str) and dst.value.startswith(FORMULA_PREFIXES):
        dst.data_type = "s"
        dst.quotePrefix = True


def latinize(text):
    if not isinstance(text, str):
        return text
    out = []
    chars = list(text)
    for idx, char in enumerate(chars):
        mapped = SR_CYR_TO_LAT.get(char)
        if mapped is None:
            out.append(char)
            continue
        if char in DIGRAPH_UPPER:
            prev_char = chars[idx - 1] if idx > 0 else ""
            next_char = chars[idx + 1] if idx + 1 < len(chars) else ""
            if (prev_char in CYR_UPPER or not (prev_char in CYR_LOWER)) and next_char in CYR_UPPER:
                mapped = DIGRAPH_UPPER[char]
        out.append(mapped)
    return "".join(out)


def backup_existing_files():
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backed_up = []
    for path in [SOURCE] + [UPLOAD / f"sr{i}.xlsx" for i in range(1, 8)]:
        if path.exists():
            dst = BACKUP_DIR / path.name
            if not dst.exists():
                shutil.copy2(path, dst)
            backed_up.append(str(dst))
    return backed_up


def latinize_source_workbook():
    wb = openpyxl.load_workbook(SOURCE, data_only=False)
    ws = wb.active
    if ws.cell(2, 3).value != "sr":
        raise ValueError(f"sr.xlsx 第 3 列表头异常: {ws.cell(2, 3).value!r}")

    changed_rows = 0
    changed_cells = 0
    cyr_before = 0
    samples = []

    for row_idx in range(3, ws.max_row + 1):
        cell = ws.cell(row_idx, 3)
        old = cell.value
        if not isinstance(old, str):
            continue
        if CYRILLIC_RE.search(old):
            cyr_before += 1
        new = latinize(old)
        if new != old:
            set_text(cell, new)
            changed_rows += 1
            changed_cells += 1
            if len(samples) < 30:
                samples.append(
                    {
                        "row": row_idx,
                        "keyId": ws.cell(row_idx, 1).value,
                        "before": old[:160],
                        "after": new[:160],
                    }
                )

    wb.save(SOURCE)
    return {
        "path": str(SOURCE),
        "rows": ws.max_row - 2,
        "cyrillic_rows_before": cyr_before,
        "changed_rows": changed_rows,
        "changed_cells": changed_cells,
        "samples": samples,
    }


def write_part(src_ws, out_path, start_row, end_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sr"
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")

    for row_idx in (1, 2):
        for col_idx in range(1, 5):
            copy_cell(src_ws.cell(row_idx, col_idx), ws.cell(row_idx, col_idx))
        ws.row_dimensions[row_idx].height = src_ws.row_dimensions[row_idx].height

    out_row = 3
    for src_row in range(start_row, end_row + 1):
        for col_idx in range(1, 5):
            copy_cell(src_ws.cell(src_row, col_idx), ws.cell(out_row, col_idx))
        out_row += 1

    for col in ("A", "B", "C", "D"):
        ws.column_dimensions[col].width = src_ws.column_dimensions[col].width

    wb.save(out_path)


def split_source_workbook():
    wb = openpyxl.load_workbook(SOURCE, data_only=False)
    ws = wb.active
    data_rows = ws.max_row - 2
    outputs = []
    part_no = 1
    for start in range(3, ws.max_row + 1, MAX_DATA_ROWS):
        end = min(ws.max_row, start + MAX_DATA_ROWS - 1)
        out_path = UPLOAD / f"sr{part_no}.xlsx"
        write_part(ws, out_path, start, end)
        outputs.append({"path": str(out_path), "rows": end - start + 1})
        part_no += 1
    if sum(item["rows"] for item in outputs) != data_rows:
        raise ValueError("拆分后总行数与源文件不一致")
    return outputs


def verify_files(paths):
    all_keys = set()
    duplicate_keys = []
    files = []
    cyr_samples = []
    formula_cells = []
    formula_prefix_without_quote = []
    empty_target_unmarked = []
    total_rows = 0
    marked_rows = 0
    cyrillic_rows = 0

    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=False)
        ws = wb.active
        headers = [ws.cell(2, col).value for col in range(1, 5)]
        if headers != ["key ID", "en", "sr", "备注"]:
            raise ValueError(f"{path.name} 表头异常: {headers}")

        file_cyr = 0
        file_marked = 0
        for row_idx in range(3, ws.max_row + 1):
            total_rows += 1
            key = str(ws.cell(row_idx, 1).value or "").strip()
            target = "" if ws.cell(row_idx, 3).value is None else str(ws.cell(row_idx, 3).value)
            remark = str(ws.cell(row_idx, 4).value or "").strip()
            if key in all_keys:
                duplicate_keys.append(key)
            all_keys.add(key)
            if remark:
                marked_rows += 1
                file_marked += 1
            if not target.strip() and not remark:
                empty_target_unmarked.append(key)
            if CYRILLIC_RE.search(target):
                cyrillic_rows += 1
                file_cyr += 1
                if len(cyr_samples) < 20:
                    cyr_samples.append({"file": str(path), "row": row_idx, "keyId": key, "value": target[:160]})
            for col_idx in range(1, 5):
                cell = ws.cell(row_idx, col_idx)
                if cell.data_type == "f":
                    formula_cells.append(f"{path.name}!{cell.coordinate}")
                if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES) and not cell.quotePrefix:
                    formula_prefix_without_quote.append(f"{path.name}!{cell.coordinate}")

        files.append(
            {
                "file": str(path),
                "rows": ws.max_row - 2,
                "marked_rows": file_marked,
                "cyrillic_rows": file_cyr,
            }
        )

    issues = []
    if duplicate_keys:
        issues.append(f"DUPLICATE_KEYS count={len(duplicate_keys)} sample={duplicate_keys[:10]}")
    if cyrillic_rows:
        issues.append(f"CYRILLIC_REMAINING count={cyrillic_rows} sample={cyr_samples[:5]}")
    if empty_target_unmarked:
        issues.append(f"EMPTY_TARGET_UNMARKED count={len(empty_target_unmarked)} sample={empty_target_unmarked[:10]}")
    if formula_cells:
        issues.append(f"FORMULA_CELLS count={len(formula_cells)} sample={formula_cells[:10]}")
    if formula_prefix_without_quote:
        issues.append(
            "FORMULA_PREFIX_WITHOUT_QUOTE "
            f"count={len(formula_prefix_without_quote)} sample={formula_prefix_without_quote[:10]}"
        )

    return {
        "files": files,
        "combined_rows": total_rows,
        "unique_key_count": len(all_keys),
        "duplicate_key_count": len(duplicate_keys),
        "empty_target_unmarked_count": len(empty_target_unmarked),
        "formula_cell_count": len(formula_cells),
        "formula_prefix_without_quote_count": len(formula_prefix_without_quote),
        "marked_rows": marked_rows,
        "cyrillic_rows": cyrillic_rows,
        "cyrillic_samples": cyr_samples,
        "pass": not issues,
        "issues": issues,
    }


def main():
    backed_up = backup_existing_files()
    latinize_summary = latinize_source_workbook()
    parts = split_source_workbook()
    part_paths = [Path(item["path"]) for item in parts]
    combined_qc = verify_files([SOURCE])
    split_qc = verify_files(part_paths)

    report = {
        "lang": "sr",
        "run_date": datetime.now().strftime("%Y%m%d"),
        "mode": "latinize_serbian_upload",
        "backups": backed_up,
        "latinize_summary": latinize_summary,
        "split_outputs": parts,
        "combined_qc": combined_qc,
        "split_qc": split_qc,
        "pass": combined_qc["pass"] and split_qc["pass"],
    }
    report_path = PROCESS / "sr_latin_qc_20260622.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
