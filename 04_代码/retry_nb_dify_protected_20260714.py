#!/usr/bin/env python3
"""对 nb Dify 未解决项做 token 保护后二次重试，并汇总 v2 上传文件。"""

import argparse
import hashlib
import json
import os
import re
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
BASE_FINAL = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_最终可上传_20260714.xlsx"
UNRESOLVED = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_仍需处理_20260714.xlsx"
JSONL = BASE / "02_过程数据/nb_dify_protected_retry_20260714.jsonl"
REPORT = BASE / "02_过程数据/nb_dify_protected_retry_report_20260714.json"
FINAL_V2 = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_最终可上传_v2_20260714.xlsx"
UNRESOLVED_V2 = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_仍需处理_v2_20260714.xlsx"

DIFY_PROXY = os.environ.get("DIFY_PROXY") or "http://127.0.0.1:30001/rest"
DIFY_PATH = "/aihelper-dify-api2/v1/workflows/run"
PH_RE = re.compile(r"(\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+)")
TAG_RE = re.compile(r"<[^>]+>")
FORMULA_PREFIXES = ("=", "+", "-", "@")
write_lock = threading.Lock()

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)


def norm(value):
    return "" if value is None else str(value)


def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", norm(value))


def md5(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def placeholders(text):
    return PH_RE.findall(norm(text))


def tags(text):
    return TAG_RE.findall(norm(text))


def validate(en, nb):
    issues = []
    if not nb.strip():
        issues.append("EMPTY_NB")
    if sorted(placeholders(en)) != sorted(placeholders(nb)):
        issues.append("PH_MISMATCH")
    if tags(en) and tags(en) != tags(nb):
        issues.append("HTML_TAG_MISMATCH")
    return issues


def protect(text):
    mapping = []

    def repl_tag(match):
        token = f"ZXSHEINTAG{len(mapping)}ZX"
        mapping.append((token, match.group(0)))
        return token

    def repl_ph(match):
        token = f"ZXSHEINPH{len(mapping)}ZX"
        mapping.append((token, match.group(0)))
        return token

    safe = TAG_RE.sub(repl_tag, text)
    safe = PH_RE.sub(repl_ph, safe)
    return safe, mapping


def restore(text, mapping):
    restored = text
    for token, raw in mapping:
        restored = restored.replace(token, raw)
    return restored


def set_text(cell, value):
    text = clean(value)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"
    if text.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def load_workbook_rows(path, upload_format):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_row = rows[1] if upload_format else rows[0]
    headers = [norm(value).strip() for value in header_row]
    body = rows[2:] if upload_format else rows[1:]
    key_idx = headers.index("key ID")
    en_idx = headers.index("en")
    nb_idx = headers.index("nb")
    result = []
    for row in body:
        key = norm(row[key_idx]).strip()
        if not key:
            continue
        result.append({
            "keyId": key,
            "en": norm(row[en_idx]).strip(),
            "nb": norm(row[nb_idx]).strip(),
        })
    return result


def load_done():
    done = {}
    if not JSONL.exists():
        return done
    with open(JSONL, encoding="utf-8") as f:
        for line in f:
            try:
                record = json.loads(line)
            except Exception:
                continue
            if record.get("en_hash"):
                done[record["en_hash"]] = record
    return done


def append_jsonl(record):
    with write_lock:
        with open(JSONL, "a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def extract_translation(resp):
    outputs = (resp.get("data") or {}).get("outputs") or {}
    output = outputs.get("output")
    if isinstance(output, list) and output:
        return norm(output[0].get("translated_text")).strip(), output[0]
    if isinstance(output, str):
        return output.strip(), {"output": output}
    return "", outputs


def call_dify(en, model, api_key, timeout):
    safe, mapping = protect(en)
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "mcp-proxy-path": DIFY_PATH,
        "mcp-proxy-env": "prod",
        "mcp-proxy-dc": "aliyun_dc",
    }
    payload = {
        "inputs": {
            "model": model,
            "OriginalLanguage": "English",
            "OriginalText": safe,
            "TargetLanguage": "nb",
            "TextContext": "Translate to Norwegian Bokmål. Keep all tokens like ZXSHEINTAG0ZX and ZXSHEINPH0ZX exactly unchanged.",
        },
        "response_mode": "blocking",
        "user": "codex-nb-protected-retry",
    }
    response = requests.post(DIFY_PROXY, headers=headers, data=json.dumps(payload, ensure_ascii=False), timeout=timeout)
    response.raise_for_status()
    text, raw = extract_translation(response.json())
    return restore(text, mapping), raw, safe


def translate_one(en, api_key, primary_model, fallback_model, timeout):
    attempts = []
    for model in [primary_model, fallback_model]:
        try:
            nb, raw, safe = call_dify(en, model, api_key, timeout)
            issues = validate(en, nb)
            attempts.append({"model": model, "issues": issues, "nb": nb, "safe": safe[:500], "raw": raw})
            if not issues:
                return {"en_hash": md5(en), "en": en, "nb": nb, "issues": [], "attempts": attempts, "status": "ok", "ts": datetime.now(timezone.utc).isoformat()}
        except Exception as exc:
            attempts.append({"model": model, "error": str(exc)})
    best = attempts[-1] if attempts else {}
    return {"en_hash": md5(en), "en": en, "nb": best.get("nb", ""), "issues": best.get("issues", ["DIFY_ERROR"]), "attempts": attempts, "status": "ok" if best.get("nb") else "error", "ts": datetime.now(timezone.utc).isoformat()}


def write_upload(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "nb"
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165
    for col, header in enumerate(["key ID", "en", "nb", "备注"], 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 3):
        for col_idx, value in enumerate([row["keyId"], row["en"], row["nb"], ""], 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 28, "B": 70, "C": 70, "D": 20}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def write_unresolved(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "仍需处理"
    headers = ["key ID", "en", "nb", "问题类型"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F4CCCC")
    for row_idx, row in enumerate(rows, 2):
        values = [row["keyId"], row["en"], row.get("nb", ""), ";".join(row.get("issues", []))]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 28, "B": 70, "C": 70, "D": 28}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--timeout", type=int, default=180)
    parser.add_argument("--primary-model", default="gemini-3-flash")
    parser.add_argument("--fallback-model", default="gpt-5-mini")
    args = parser.parse_args()
    api_key = os.environ.get("DIFY_NB_API_KEY")
    if not api_key:
        raise RuntimeError("缺少环境变量 DIFY_NB_API_KEY")

    base_rows = load_workbook_rows(BASE_FINAL, upload_format=True)
    unresolved_rows = load_workbook_rows(UNRESOLVED, upload_format=False)
    unique = {md5(row["en"]): row["en"] for row in unresolved_rows}
    done = load_done()
    pending = [(h, en) for h, en in unique.items() if h not in done]
    pending.sort(key=lambda item: (len(item[1]), item[1]))
    print(f"base={len(base_rows)} unresolved={len(unresolved_rows)} unique={len(unique)} pending={len(pending)}", flush=True)

    if pending:
        counters = {"ok": 0, "error": 0}
        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {executor.submit(translate_one, en, api_key, args.primary_model, args.fallback_model, args.timeout): (h, en) for h, en in pending}
            for idx, future in enumerate(as_completed(futures), 1):
                h, en = futures[future]
                try:
                    record = future.result()
                except Exception as exc:
                    record = {"en_hash": h, "en": en, "nb": "", "issues": ["DIFY_EXCEPTION"], "status": "error", "error": str(exc), "ts": datetime.now(timezone.utc).isoformat()}
                append_jsonl(record)
                done[record["en_hash"]] = record
                if record.get("status") == "ok" and not record.get("issues"):
                    counters["ok"] += 1
                else:
                    counters["error"] += 1
                if idx <= 5 or idx % 20 == 0:
                    print(f"进度 {idx}/{len(pending)} ok={counters['ok']} error={counters['error']}", flush=True)

    done = load_done()
    final_rows = list(base_rows)
    still_bad = []
    for row in unresolved_rows:
        record = done.get(md5(row["en"]), {})
        nb = record.get("nb", "")
        issues = validate(row["en"], nb)
        if nb and not issues:
            final_rows.append({"keyId": row["keyId"], "en": row["en"], "nb": nb})
        else:
            still_bad.append({"keyId": row["keyId"], "en": row["en"], "nb": nb, "issues": issues or record.get("issues") or ["NO_TRANSLATION"]})

    write_upload(FINAL_V2, final_rows)
    write_unresolved(UNRESOLVED_V2, still_bad)
    reason_counts = {}
    for row in still_bad:
        reason = ";".join(row["issues"])
        reason_counts[reason] = reason_counts.get(reason, 0) + 1
    report = {
        "base_rows": len(base_rows),
        "protected_retry_rows": len(unresolved_rows),
        "final_rows": len(final_rows),
        "still_bad_rows": len(still_bad),
        "still_bad_reason_counts": reason_counts,
        "jsonl": str(JSONL),
        "final_file": str(FINAL_V2),
        "still_bad_file": str(UNRESOLVED_V2),
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
