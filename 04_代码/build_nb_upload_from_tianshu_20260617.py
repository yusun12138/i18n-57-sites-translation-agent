#!/usr/bin/env python3
"""
把天枢翻译结果按 en(text) 对回原始 key，生成文案系统上传格式。

天枢输出只有：
  text: 原英文
  nb: 翻译后的书面挪威语
  nb-channel: 翻译渠道（忽略）

原始 key 映射来自拆分任务缓存：
  02_过程数据/nb_full_20260616_bokmal_dify_20260617_tasks.json
"""

import json
import os
import re
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
TASK_JSON = BASE / "02_过程数据" / "nb_full_20260616_bokmal_dify_20260617_tasks.json"
OUTPUT_FILES = [
    Path("/Users/10304327/Downloads/nb_tianshu_001_output.xlsx"),
    Path("/Users/10304327/Downloads/nb_tianshu_002_output.xlsx"),
]

PROCESS = BASE / "02_过程数据"
UPLOAD = BASE / "05_上传格式"
JSONL_OUT = PROCESS / "nb_full_20260616_tianshu_byteplus_20260617.jsonl"
QC_OUT = PROCESS / "nb_tianshu_byteplus_20260617_upload_qc.json"
XLSX_OUT = UPLOAD / "nb_tianshu_byteplus_20260617.xlsx"

LANG = "nb"
FORMULA_PREFIXES = ("=", "+", "-", "@")
PH_RE = re.compile(r"\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+")
TAG_RE = re.compile(r"<[^>]+>")
ESC_RE = re.compile(r"\\['\u2018\u2019]")

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RISK_FILL = PatternFill("solid", fgColor="FFF2CC")

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)


def norm(value):
    return "" if value is None else str(value).strip()


def qc_issues(en, tr):
    en = str(en or "")
    tr = str(tr or "")
    issues = []
    if not tr.strip():
        issues.append("EMPTY")
    if sorted(PH_RE.findall(en)) != sorted(PH_RE.findall(tr)):
        issues.append("PH_MISMATCH")
    if ESC_RE.findall(en) != ESC_RE.findall(tr):
        issues.append("ESCAPED_SINGLE_QUOTE_MISMATCH")
    if TAG_RE.findall(en) != TAG_RE.findall(tr):
        issues.append("TAG_SEQUENCE_MISMATCH")
    if len(en) >= 1200 and len(tr) < len(en) * 0.45:
        issues.append(f"POSSIBLE_TRUNCATED_LONG_TEXT en_len={len(en)} got_len={len(tr)}")
    return issues


def load_missing_rows():
    data = json.loads(TASK_JSON.read_text(encoding="utf-8"))
    missing_rows = OrderedDict(data["missing_rows"])
    needed = [str(text) for text in data["needed"] if str(text).strip()]
    return missing_rows, needed


def load_tianshu_translations():
    translations = {}
    duplicate_same = 0
    conflicts = []
    file_stats = []

    for path in OUTPUT_FILES:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [norm(v) for v in next(rows)]
        try:
            text_i = headers.index("text")
            nb_i = headers.index("nb")
        except ValueError as exc:
            raise ValueError(f"{path.name} 表头必须包含 text 和 nb，实际: {headers}") from exc

        data_rows = 0
        valid_rows = 0
        empty_nb = 0
        for row in rows:
            data_rows += 1
            en = norm(row[text_i] if text_i < len(row) else "")
            nb = norm(row[nb_i] if nb_i < len(row) else "")
            if not en:
                continue
            if not nb:
                empty_nb += 1
                continue
            valid_rows += 1
            if en in translations:
                if translations[en] == nb:
                    duplicate_same += 1
                else:
                    conflicts.append({"en": en, "first": translations[en], "second": nb, "file": str(path)})
                continue
            translations[en] = nb

        file_stats.append({
            "file": str(path),
            "headers": headers,
            "data_rows": data_rows,
            "valid_rows": valid_rows,
            "empty_nb": empty_nb,
        })

    return translations, duplicate_same, conflicts, file_stats


def set_text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    if cell.value.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def write_jsonl(translations, needed):
    needed_set = set(needed)
    with open(JSONL_OUT, "w", encoding="utf-8") as f:
        for en in needed:
            nb = translations.get(en, "")
            record = {
                "en_full": en,
                "translation": nb,
                "status": "ok" if nb and not qc_issues(en, nb) else "blocked",
                "engine": "tianshu_byteplus",
                "target_language": LANG,
                "issues": qc_issues(en, nb),
            }
            if not record["issues"]:
                record.pop("issues")
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
    return len(needed_set)


def build_upload(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = LANG
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165

    for col, header in enumerate(["key ID", "en", LANG, "备注"], 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, item in enumerate(rows, 3):
        for col_idx, value in enumerate([item["key"], item["en"], item["nb"], item["remark"]], 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if item["issues"]:
            ws.cell(row_idx, 3).fill = RISK_FILL
            ws.cell(row_idx, 4).fill = RISK_FILL

    for col, width in {"A": 28, "B": 62, "C": 62, "D": 46}.items():
        ws.column_dimensions[col].width = width

    UPLOAD.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)


def verify_workbook(expected_rows):
    wb = openpyxl.load_workbook(XLSX_OUT, read_only=True, data_only=False)
    ws = wb.active
    headers = [norm(ws.cell(row=2, column=i).value) for i in range(1, 5)]
    if headers != ["key ID", "en", LANG, "备注"]:
        raise ValueError(f"上传文件表头异常: {headers}")
    actual_rows = ws.max_row - 2
    if actual_rows != expected_rows:
        raise ValueError(f"上传文件行数异常: expected={expected_rows}, actual={actual_rows}")


def main():
    missing_rows, needed = load_missing_rows()
    translations, duplicate_same, conflicts, file_stats = load_tianshu_translations()
    if conflicts:
        raise ValueError(f"同一 text 出现不同 nb 翻译，请先处理: {conflicts[:5]}")

    missing_texts = sorted(set(needed) - set(translations))
    extra_texts = sorted(set(translations) - set(needed))
    write_jsonl(translations, needed)

    rows = []
    issue_counter = Counter()
    source_counter = Counter()
    for item in missing_rows.values():
        en = item["en"]
        nb = translations.get(en, "")
        issues = qc_issues(en, nb)
        for issue in issues:
            issue_counter[issue.split()[0]] += 1
        if not nb:
            source_counter["missing_translation"] += 1
        else:
            source_counter["tianshu"] += 1
        remark = ""
        if issues:
            remark = f"质检风险：{';'.join(issues)}；请确认后再上传"
        rows.append({"key": item["key"], "en": en, "nb": nb, "issues": issues, "remark": remark})

    build_upload(rows)
    verify_workbook(len(rows))

    qc = {
        "output_xlsx": str(XLSX_OUT),
        "jsonl": str(JSONL_OUT),
        "source_files": file_stats,
        "unique_needed": len(set(needed)),
        "unique_translated": len(translations),
        "duplicate_same": duplicate_same,
        "missing_unique_text_count": len(missing_texts),
        "missing_unique_text_sample": missing_texts[:50],
        "extra_unique_text_count": len(extra_texts),
        "extra_unique_text_sample": extra_texts[:20],
        "expanded_key_rows": len(rows),
        "source_counter": dict(source_counter),
        "issue_counter": dict(issue_counter),
        "pass": len(rows) == len(missing_rows) and len(missing_texts) == 0,
        "risk_row_count": sum(1 for row in rows if row["issues"]),
        "risk_samples": [
            {
                "key": row["key"],
                "en": row["en"][:160],
                "nb": row["nb"][:160],
                "issues": row["issues"],
            }
            for row in rows
            if row["issues"]
        ][:50],
    }
    QC_OUT.write_text(json.dumps(qc, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(qc, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
