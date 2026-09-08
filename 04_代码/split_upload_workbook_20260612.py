#!/usr/bin/env python3
import argparse
import os
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

BASE = Path(os.environ.get('I18N_WORKDIR') or Path(__file__).resolve().parents[1]).resolve()
UPLOAD = BASE / '05_上传格式'
FORMULA_PREFIXES = ('=', '+', '-', '@')


def copy_cell(src, dst):
    if isinstance(src, MergedCell):
        return
    dst.value = src.value
    dst.data_type = src.data_type
    dst.number_format = src.number_format
    dst.quotePrefix = src.quotePrefix
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    if isinstance(dst.value, str) and dst.value.startswith(FORMULA_PREFIXES):
        dst.data_type = 's'
        dst.quotePrefix = True


def write_part(src_ws, out_path, start_row, end_row):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = src_ws.title
    ws.freeze_panes = 'A3'
    ws.merge_cells('A1:D1')

    for r in (1, 2):
        for c in range(1, 5):
            copy_cell(src_ws.cell(r, c), ws.cell(r, c))
        ws.row_dimensions[r].height = src_ws.row_dimensions[r].height

    out_r = 3
    for src_r in range(start_row, end_row + 1):
        for c in range(1, 5):
            copy_cell(src_ws.cell(src_r, c), ws.cell(out_r, c))
        out_r += 1

    for col in ('A', 'B', 'C', 'D'):
        ws.column_dimensions[col].width = src_ws.column_dimensions[col].width
    wb.save(out_path)


def verify(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    formulas = []
    prefix_without_quote = []
    keys = set()
    dupes = []
    for r in range(3, ws.max_row + 1):
        key = ws.cell(r, 1).value
        if key in keys:
            dupes.append(key)
        keys.add(key)
        for c in range(1, 5):
            cell = ws.cell(r, c)
            if cell.data_type == 'f':
                formulas.append(cell.coordinate)
            if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES) and not cell.quotePrefix:
                prefix_without_quote.append(cell.coordinate)
    return {
        'file': str(path),
        'rows': ws.max_row - 2,
        'duplicate_keys': len(dupes),
        'formula_cells': len(formulas),
        'formula_prefix_without_quote': len(prefix_without_quote),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--lang', required=True, choices=['pl', 'nl', 'el-gr', 'ru', 'nb', 'sl', 'lv', 'sr', 'et'])
    parser.add_argument('--parts', type=int, default=2)
    args = parser.parse_args()

    source = UPLOAD / f'{args.lang}.xlsx'
    wb = openpyxl.load_workbook(source, data_only=False)
    ws = wb.active
    data_rows = ws.max_row - 2
    part_size = (data_rows + args.parts - 1) // args.parts
    outputs = []

    for idx in range(args.parts):
        start = 3 + idx * part_size
        end = min(2 + data_rows, start + part_size - 1)
        if start > end:
            continue
        out = UPLOAD / f'{args.lang}{idx + 1}.xlsx'
        write_part(ws, out, start, end)
        outputs.append(verify(out))

    print({'source': str(source), 'source_rows': data_rows, 'outputs': outputs})


if __name__ == '__main__':
    main()
