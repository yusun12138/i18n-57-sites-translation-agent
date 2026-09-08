#!/usr/bin/env python3
"""从最新全量导出文件中提取 keyId + en 清单。"""

import csv
import hashlib
import json
import os
import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
RAW = BASE / "01_原始数据"
PROCESS = BASE / "02_过程数据"
RUN_DATE = "20260707"


def norm(value):
    return "" if value is None else str(value).strip()


def en_hash(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def latest(pattern):
    files = sorted(RAW.glob(pattern), key=lambda path: path.stat().st_mtime)
    if not files:
        raise FileNotFoundError(pattern)
    return files[-1]


def source_files():
    files = []
    for platform in [1, 2, 3, 4, 5]:
        path = latest(f"type1_前端key_platform{platform}_full_{RUN_DATE}_*.xlsx")
        files.append(
            {
                "path": path,
                "type": "type1_前端key",
                "platform": f"platform{platform}",
            }
        )
    files.append(
        {
            "path": latest(f"type2_提示语_full_{RUN_DATE}_*.xlsx"),
            "type": "type2_提示语",
            "platform": "platform6_7",
        }
    )
    files.append(
        {
            "path": latest(f"type3_错误码_full_{RUN_DATE}_*.xlsx"),
            "type": "type3_错误码",
            "platform": "platform12",
        }
    )
    return files


def write_xlsx(path, headers, rows, sheet_name):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    wb.save(path)


def main():
    PROCESS.mkdir(parents=True, exist_ok=True)
    rows = []
    unique = OrderedDict()
    file_stats = []

    for item in source_files():
        path = item["path"]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()
        iterator = ws.iter_rows(values_only=True)
        headers = [norm(value) for value in next(iterator)]
        lower_headers = [header.lower() for header in headers]
        key_idx = next(i for i, header in enumerate(headers) if header in {"keyId", "key ID", "key_id"})
        en_idx = lower_headers.index("en")

        total_rows = 0
        rows_with_en = 0
        for source_row_idx, row in enumerate(iterator, 2):
            total_rows += 1
            key = norm(row[key_idx] if key_idx < len(row) else "")
            en = norm(row[en_idx] if en_idx < len(row) else "")
            if not key or not en:
                continue
            rows_with_en += 1
            digest = en_hash(en)
            record = {
                "source_file": path.name,
                "type": item["type"],
                "platform": item["platform"],
                "source_row": source_row_idx,
                "keyId": key,
                "en": en,
                "en_hash": digest,
            }
            rows.append(record)
            if digest not in unique:
                unique[digest] = {
                    "en_hash": digest,
                    "en": en,
                    "key_count": 0,
                    "sample_keyIds": [],
                    "sample_sources": [],
                }
            unique_item = unique[digest]
            unique_item["key_count"] += 1
            if len(unique_item["sample_keyIds"]) < 20:
                unique_item["sample_keyIds"].append(key)
            if len(unique_item["sample_sources"]) < 10:
                unique_item["sample_sources"].append(f"{path.name}:{source_row_idx}")

        file_stats.append(
            {
                "file": path.name,
                "type": item["type"],
                "platform": item["platform"],
                "total_rows": total_rows,
                "rows_with_en": rows_with_en,
            }
        )

    full_headers = ["source_file", "type", "platform", "source_row", "keyId", "en", "en_hash"]
    unique_headers = ["en_hash", "en", "key_count", "sample_keyIds", "sample_sources"]
    unique_rows = []
    for item in unique.values():
        unique_rows.append(
            {
                "en_hash": item["en_hash"],
                "en": item["en"],
                "key_count": item["key_count"],
                "sample_keyIds": "|".join(item["sample_keyIds"]),
                "sample_sources": "|".join(item["sample_sources"]),
            }
        )

    full_csv = PROCESS / f"full_en_{RUN_DATE}.csv"
    with open(full_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=full_headers)
        writer.writeheader()
        writer.writerows(rows)

    unique_csv = PROCESS / f"full_en_unique_{RUN_DATE}.csv"
    with open(unique_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=unique_headers)
        writer.writeheader()
        writer.writerows(unique_rows)

    full_xlsx = PROCESS / f"full_en_{RUN_DATE}.xlsx"
    write_xlsx(full_xlsx, full_headers, rows, "full_en")

    unique_xlsx = PROCESS / f"full_en_unique_{RUN_DATE}.xlsx"
    write_xlsx(unique_xlsx, unique_headers, unique_rows, "unique_en")

    duplicate_en_count = sum(1 for item in unique.values() if item["key_count"] > 1)
    report = {
        "run_date": RUN_DATE,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "input_files": file_stats,
        "summary": {
            "full_rows": len(rows),
            "unique_en": len(unique_rows),
            "duplicate_en_groups": duplicate_en_count,
        },
        "outputs": {
            "full_csv": str(full_csv),
            "full_xlsx": str(full_xlsx),
            "unique_csv": str(unique_csv),
            "unique_xlsx": str(unique_xlsx),
        },
    }
    report_path = PROCESS / f"full_en_report_{RUN_DATE}.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
