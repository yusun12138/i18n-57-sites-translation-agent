#!/usr/bin/env python3
"""重翻 nb/sr 当前平台全量中仍等于英文的内容。

凭据通过环境变量 DIFY_API_KEY 或 DIFY_NB_API_KEY 传入，不写入脚本。
"""

import argparse
import hashlib
import json
import os
import re
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
DATA_DIR = BASE / "01_原始数据"
PROCESS_DIR = BASE / "02_过程数据" / "nb_sr_english_residual_20260716"
REPORT_DIR = BASE / "03_分析报告"
OUTPUT_DIR = BASE / "05_上传格式" / "20260716_nb_sr_英文残留重翻"
MAX_UPLOAD_ROWS = 25000

DIFY_PROXY = os.environ.get("DIFY_PROXY") or "http://127.0.0.1:30001/rest"
DIFY_PATH = "/aihelper-dify-api2/v1/workflows/run"
DEFAULT_MODEL = "gemini-3-flash"
FALLBACK_MODEL = "gpt-5-mini"

LANGS = ("nb", "sr")
TARGET_NAME = {
    "nb": "Norwegian Bokmål (书面挪威语)",
    "sr": "Serbian in Latin script only (塞尔维亚语拉丁字母)",
}

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)

PH_RE = re.compile(r"(\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+|\$\{[^}]+\})")
TAG_RE = re.compile(r"<[^>]+>")
ENTITY_RE = re.compile(r"&[a-zA-Z#0-9]+;")
URL_RE = re.compile(
    r"(https?://|www\.|\.com\b|\.net\b|\.cn\b|\.jpg\b|\.jpeg\b|\.png\b|"
    r"\.webp\b|\.gif\b|\.svg\b|\.mp4\b|/assets?/|^//)",
    re.I,
)
NON_LATIN_SOURCE_RE = re.compile(r"[\u4e00-\u9fff\u3040-\u30ff\u0400-\u04ff\u0600-\u06ff]")
CYRILLIC_RE = re.compile(r"[\u0400-\u04FF]")
FORMULA_PREFIXES = ("=", "+", "-", "@")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RISK_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="F4CCCC")

BRAND_ONLY = {
    "shein", "romwe", "motf", "dazy", "evolushein", "sheglam", "shein.com",
    "facebook", "instagram", "tiktok", "youtube", "google", "apple", "paypal",
    "klarna", "afterpay", "tabby", "tamara", "dji", "nintendo",
}
SAFE_CODE_WORDS = {
    "sku", "id", "uid", "url", "uri", "api", "pc", "app", "ios", "android",
    "dji", "gb", "mb", "xl", "xxl", "vip", "cod", "cvv", "iban", "bic",
}
ENGLISH_UI_WORDS = {
    "about", "accept", "account", "add", "address", "apply", "available",
    "back", "bag", "balance", "buy", "cancel", "card", "cart", "cash",
    "checkout", "code", "color", "complete", "confirm", "congrats", "continue",
    "coupon", "credit", "currency", "custom", "deal", "delete", "delivery",
    "details", "discount", "done", "edit", "email", "error", "expired",
    "failed", "fee", "filter", "free", "get", "gift", "help", "history",
    "instant", "invite", "item", "items", "login", "mark", "method", "more",
    "new", "next", "order", "paid", "password", "pay", "payment", "phone",
    "points", "price", "privacy", "processing", "refund", "remove", "return",
    "review", "reward", "save", "search", "select", "settings", "ship",
    "shipping", "shop", "sign", "size", "status", "submit", "success",
    "successful", "terms", "top", "total", "track", "try", "update", "user",
    "valued", "verification", "view", "wallet",
}

SR_CYR_TO_LAT = {
    "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Ђ": "Đ", "Е": "E", "Ж": "Ž",
    "З": "Z", "И": "I", "Ј": "J", "К": "K", "Л": "L", "Љ": "Lj", "М": "M", "Н": "N",
    "Њ": "Nj", "О": "O", "П": "P", "Р": "R", "С": "S", "Т": "T", "Ћ": "Ć", "У": "U",
    "Ф": "F", "Х": "H", "Ц": "C", "Ч": "Č", "Џ": "Dž", "Ш": "Š",
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "ђ": "đ", "е": "e", "ж": "ž",
    "з": "z", "и": "i", "ј": "j", "к": "k", "л": "l", "љ": "lj", "м": "m", "н": "n",
    "њ": "nj", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "ћ": "ć", "у": "u",
    "ф": "f", "х": "h", "ц": "c", "ч": "č", "џ": "dž", "ш": "š",
    "Ё": "Jo", "Й": "J", "Щ": "Šč", "Ъ": "", "Ы": "Y", "Ь": "", "Э": "E", "Ю": "Ju", "Я": "Ja",
    "ё": "jo", "й": "j", "щ": "šč", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "ju", "я": "ja",
}

write_lock = threading.Lock()


def log(message):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def norm(value):
    text = "" if value is None else str(value)
    text = text.replace("\\n", "\n").replace("\\r", "\n")
    return text.strip()


def clean(value):
    text = "" if value is None else str(value)
    text = text.replace("\\n", "\n").replace("\\r", "\n")
    return ILLEGAL_CHARACTERS_RE.sub("", text).strip()


def md5(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def latest_export_files(timestamp=None):
    if timestamp:
        files = sorted(DATA_DIR.glob(f"*full_{timestamp}.xlsx"))
    else:
        groups = defaultdict(list)
        for path in DATA_DIR.glob("*full_*.xlsx"):
            match = re.search(r"_full_(\d{8}_\d{6})\.xlsx$", path.name)
            if match:
                groups[match.group(1)].append(path)
        if not groups:
            raise FileNotFoundError(f"没有找到全量导出文件: {DATA_DIR}")
        timestamp = sorted(groups)[-1]
        files = sorted(groups[timestamp])
    if not files:
        raise FileNotFoundError(f"没有找到 timestamp={timestamp} 的全量导出文件")
    return timestamp, files


def strip_structural(text):
    text = TAG_RE.sub("", text)
    text = PH_RE.sub("", text)
    text = ENTITY_RE.sub("", text)
    return text.strip()


def classify_equal_en(en):
    text = norm(en)
    if not text:
        return "EMPTY_EN", False
    if URL_RE.search(text):
        return "URL_OR_ASSET", False
    if NON_LATIN_SOURCE_RE.search(text):
        return "NON_LATIN_SOURCE", False

    stripped = strip_structural(text)
    stripped_no_brand = re.sub(r"\b(SHEIN|ROMWE|MOTF|DAZY|SHEGLAM)\b", "", stripped, flags=re.I).strip()
    words = re.findall(r"[A-Za-z]+", stripped_no_brand)
    lower_words = [word.lower() for word in words]

    if not words:
        return "PLACEHOLDER_OR_SYMBOL_ONLY", False
    if len(lower_words) == 1 and lower_words[0] in BRAND_ONLY:
        return "BRAND_ONLY", False

    compact = re.sub(r"[^A-Za-z0-9]", "", stripped_no_brand)
    has_digit = bool(re.search(r"\d", stripped_no_brand))
    all_caps = bool(compact) and compact.upper() == compact
    if len(compact) <= 2:
        return "CODE_OR_PRODUCT", False
    if all_caps and (has_digit or any(word.lower() in SAFE_CODE_WORDS for word in words)):
        return "CODE_OR_PRODUCT", False
    if len(lower_words) == 1 and lower_words[0] in SAFE_CODE_WORDS:
        return "CODE_OR_PRODUCT", False

    if any(word in ENGLISH_UI_WORDS for word in lower_words):
        return "SUSPICIOUS_ENGLISH", True
    if len(lower_words) >= 2:
        return "SUSPICIOUS_ENGLISH", True
    return "NEEDS_REVIEW", True


def protect(text):
    mapping = []

    def repl_tag(match):
        token = f"ZXSHEINTAG{len(mapping)}ZX"
        mapping.append((token, match.group(0)))
        return token

    def repl_ph(match):
        token = f"ZXSHEINPH{len(mapping)}ZX"
        mapping.append((token, match.group(0)))
        return token

    safe = TAG_RE.sub(repl_tag, text)
    safe = PH_RE.sub(repl_ph, safe)
    return safe, mapping


def restore(text, mapping):
    restored = text
    for token, raw in mapping:
        restored = restored.replace(token, raw)
    return restored


def latinize_sr(text):
    if not isinstance(text, str):
        return text
    return "".join(SR_CYR_TO_LAT.get(ch, ch) for ch in text)


def placeholders(text):
    return PH_RE.findall(text or "")


def tags(text):
    return TAG_RE.findall(text or "")


def validate_translation(en, translation, lang):
    issues = []
    en = en or ""
    translation = translation or ""
    if not translation.strip():
        issues.append("EMPTY")
    if norm(translation) == norm(en):
        issues.append("STILL_ENGLISH")
    if sorted(placeholders(en)) != sorted(placeholders(translation)):
        issues.append("PH_MISMATCH")
    if tags(en) and tags(en) != tags(translation):
        issues.append("TAG_SEQUENCE_MISMATCH")
    if lang == "sr" and CYRILLIC_RE.search(translation):
        issues.append("SR_CYRILLIC_REMAINING")
    if len(en) >= 1200 and len(translation) < len(en) * 0.45:
        issues.append("POSSIBLE_TRUNCATED_LONG_TEXT")
    return issues


def scan_tasks(timestamp=None):
    timestamp, files = latest_export_files(timestamp)
    PROCESS_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)

    details = {lang: [] for lang in LANGS}
    excluded = {lang: [] for lang in LANGS}
    summary = {lang: defaultdict(int) for lang in LANGS}
    seen = {lang: set() for lang in LANGS}

    for path in files:
        log(f"扫描 {path.name}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            rows = ws.iter_rows(values_only=True)
            try:
                headers = [norm(value) for value in next(rows)]
            except StopIteration:
                continue
            if "keyId" not in headers or "en" not in headers:
                continue
            key_idx = headers.index("keyId")
            en_idx = headers.index("en")
            lang_idxs = {lang: headers.index(lang) for lang in LANGS if lang in headers}
            if not lang_idxs:
                continue
            for row_num, row in enumerate(rows, start=2):
                key_id = norm(row[key_idx] if key_idx < len(row) else "")
                en = norm(row[en_idx] if en_idx < len(row) else "")
                if not key_id or not en:
                    continue
                for lang, lang_idx in lang_idxs.items():
                    current_value = norm(row[lang_idx] if lang_idx < len(row) else "")
                    if current_value != en:
                        continue
                    reason, should_translate = classify_equal_en(en)
                    summary[lang][reason] += 1
                    record = {
                        "keyId": key_id,
                        "en": en,
                        "current": current_value,
                        "lang": lang,
                        "reason": reason,
                        "source_file": path.name,
                        "sheet": ws.title,
                        "row": row_num,
                    }
                    if should_translate:
                        dedupe_key = (key_id, en)
                        if dedupe_key in seen[lang]:
                            continue
                        seen[lang].add(dedupe_key)
                        details[lang].append(record)
                    else:
                        excluded[lang].append(record)

    for lang in LANGS:
        task_path = PROCESS_DIR / f"{lang}_english_residual_tasks_20260716.json"
        task_path.write_text(json.dumps(details[lang], ensure_ascii=False, indent=2), encoding="utf-8")
        excluded_path = PROCESS_DIR / f"{lang}_english_residual_excluded_20260716.json"
        excluded_path.write_text(json.dumps(excluded[lang], ensure_ascii=False, indent=2), encoding="utf-8")

    report_path = REPORT_DIR / "nb_sr_英文残留待重翻_20260716.xlsx"
    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "summary"
    ws.append(["timestamp", timestamp])
    ws.append([])
    ws.append(["lang", "reason", "exact_same_count", "auto_translate_tasks"])
    for lang in LANGS:
        first = True
        for reason, count in sorted(summary[lang].items()):
            ws.append([lang, reason, count, len(details[lang]) if first else ""])
            first = False
    for lang in LANGS:
        ws = wb_out.create_sheet(f"{lang}_待重翻")
        headers = ["key ID", "en", "当前值", "reason", "source_file", "sheet", "row"]
        ws.append(headers)
        for item in details[lang]:
            ws.append([item["keyId"], item["en"], item["current"], item["reason"], item["source_file"], item["sheet"], item["row"]])
        for col, width in {"A": 28, "B": 70, "C": 70, "D": 24, "E": 44, "F": 18, "G": 10}.items():
            ws.column_dimensions[col].width = width
        ws = wb_out.create_sheet(f"{lang}_排除")
        ws.append(headers)
        for item in excluded[lang]:
            ws.append([item["keyId"], item["en"], item["current"], item["reason"], item["source_file"], item["sheet"], item["row"]])
        for col, width in {"A": 28, "B": 70, "C": 70, "D": 24, "E": 44, "F": 18, "G": 10}.items():
            ws.column_dimensions[col].width = width
    wb_out.save(report_path)

    summary_json = {
        "timestamp": timestamp,
        "files": [str(path) for path in files],
        "summary": {lang: dict(summary[lang]) for lang in LANGS},
        "tasks": {lang: len(details[lang]) for lang in LANGS},
        "excluded": {lang: len(excluded[lang]) for lang in LANGS},
        "report": str(report_path),
    }
    summary_path = PROCESS_DIR / "nb_sr_english_residual_task_report_20260716.json"
    summary_path.write_text(json.dumps(summary_json, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary_json


def load_tasks():
    tasks = {}
    for lang in LANGS:
        path = PROCESS_DIR / f"{lang}_english_residual_tasks_20260716.json"
        tasks[lang] = json.loads(path.read_text(encoding="utf-8"))
    return tasks


def load_done(jsonl_path):
    done = {}
    if not jsonl_path.exists():
        return done
    with open(jsonl_path, encoding="utf-8") as f:
        for line in f:
            try:
                record = json.loads(line)
            except Exception:
                continue
            if record.get("en_hash"):
                done[record["en_hash"]] = record
    return done


def record_complete(record, langs):
    if not record or record.get("status") != "ok":
        return False
    translations = record.get("translations") or {}
    for lang in langs:
        if validate_translation(record.get("en", ""), translations.get(lang, ""), lang):
            return False
    return True


def append_jsonl(path, record):
    with write_lock:
        with open(path, "a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def extract_output(resp_json):
    outputs = (resp_json.get("data") or {}).get("outputs") or {}
    output = outputs.get("output")
    items = []
    if isinstance(output, list):
        items = output
    elif isinstance(output, str):
        items = [{"target_language_code": "", "translated_text": output}]
    return items, outputs


def call_dify(en, langs, model, api_key, timeout):
    safe, mapping = protect(en)
    target_lang = ",".join(langs)
    context = (
        "Translate SHEIN e-commerce UI copy from English. "
        "If TargetLanguage contains nb, output Norwegian Bokmål (书面挪威语), not Nynorsk. "
        "If TargetLanguage contains sr, output Serbian using Latin script only, never Cyrillic. "
        "Keep brand names such as SHEIN unchanged. "
        "Keep all tokens like ZXSHEINTAG0ZX and ZXSHEINPH0ZX exactly unchanged."
    )
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "mcp-proxy-path": DIFY_PATH,
        "mcp-proxy-env": "prod",
        "mcp-proxy-dc": "aliyun_dc",
    }
    payload = {
        "inputs": {
            "model": model,
            "OriginalLanguage": "English",
            "OriginalText": safe,
            "TargetLanguage": target_lang,
            "TextContext": context,
        },
        "response_mode": "blocking",
        "user": "codex-nb-sr-english-residual",
    }
    response = requests.post(DIFY_PROXY, headers=headers, data=json.dumps(payload, ensure_ascii=False), timeout=timeout)
    response.raise_for_status()
    items, raw = extract_output(response.json())
    result = {}
    for item in items:
        code = norm(item.get("target_language_code")).lower()
        text = restore(norm(item.get("translated_text")), mapping)
        if code.startswith("nb") or code in {"no", "norwegian bokmal", "norwegian bokmål"}:
            result["nb"] = text
        elif code.startswith("sr"):
            result["sr"] = latinize_sr(text)
        elif len(langs) == 1:
            result[langs[0]] = latinize_sr(text) if langs[0] == "sr" else text
    return result, raw


def parse_batch_output(text):
    matches = list(re.finditer(r"ZXSHEINITEM(\d{4})ZX", text or ""))
    parsed = {}
    for index, match in enumerate(matches):
        item_idx = int(match.group(1))
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        value = text[start:end].strip()
        value = re.sub(r"^[\s:：\-–—]+", "", value).strip()
        parsed[item_idx] = value
    return parsed


def call_dify_batch(batch, langs, model, api_key, timeout):
    lines = []
    for idx, item in enumerate(batch):
        lines.append(f"ZXSHEINITEM{idx:04d}ZX {item['en']}")
    block = "\n".join(lines)
    safe, mapping = protect(block)
    target_lang = ",".join(langs)
    context = (
        "Translate each numbered SHEIN e-commerce UI copy item independently from English. "
        "Return exactly one output line for each input line. "
        "Keep every item marker like ZXSHEINITEM0000ZX unchanged at the beginning of its line. "
        "If TargetLanguage contains nb, output Norwegian Bokmål (书面挪威语), not Nynorsk. "
        "If TargetLanguage contains sr, output Serbian using Latin script only, never Cyrillic. "
        "Keep brand names such as SHEIN unchanged. "
        "Keep all tokens like ZXSHEINTAG0ZX and ZXSHEINPH0ZX exactly unchanged."
    )
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "mcp-proxy-path": DIFY_PATH,
        "mcp-proxy-env": "prod",
        "mcp-proxy-dc": "aliyun_dc",
    }
    payload = {
        "inputs": {
            "model": model,
            "OriginalLanguage": "English",
            "OriginalText": safe,
            "TargetLanguage": target_lang,
            "TextContext": context,
        },
        "response_mode": "blocking",
        "user": "codex-nb-sr-english-residual-batch",
    }
    response = requests.post(DIFY_PROXY, headers=headers, data=json.dumps(payload, ensure_ascii=False), timeout=timeout)
    response.raise_for_status()
    items, raw = extract_output(response.json())
    result = {lang: {} for lang in langs}
    for item in items:
        code = norm(item.get("target_language_code")).lower()
        translated = norm(item.get("translated_text"))
        lang = ""
        if code.startswith("nb") or code in {"no", "norwegian bokmal", "norwegian bokmål"}:
            lang = "nb"
        elif code.startswith("sr"):
            lang = "sr"
        elif len(langs) == 1:
            lang = langs[0]
        if not lang or lang not in result:
            continue
        parsed = parse_batch_output(translated)
        for idx, text in parsed.items():
            restored = restore(text, mapping)
            result[lang][idx] = latinize_sr(restored) if lang == "sr" else restored
    return result, raw


def translate_one(en, langs, api_key, primary_model, fallback_model, timeout):
    attempts = []
    for model in [primary_model, fallback_model]:
        if not model:
            continue
        start = time.time()
        try:
            translations, raw = call_dify(en, langs, model, api_key, timeout)
            issues_by_lang = {}
            for lang in langs:
                if lang == "sr" and translations.get(lang):
                    translations[lang] = latinize_sr(translations[lang])
                issues_by_lang[lang] = validate_translation(en, translations.get(lang, ""), lang)
            attempts.append({
                "model": model,
                "translations": translations,
                "issues_by_lang": issues_by_lang,
                "elapsed_sec": round(time.time() - start, 2),
                "raw": raw,
            })
            if all(not issues_by_lang.get(lang) for lang in langs):
                return {
                    "en_hash": md5(en),
                    "en": en,
                    "langs": langs,
                    "translations": translations,
                    "issues_by_lang": issues_by_lang,
                    "attempts": attempts,
                    "status": "ok",
                    "ts": datetime.now(timezone.utc).isoformat(),
                }
        except Exception as exc:
            attempts.append({
                "model": model,
                "error": str(exc),
                "elapsed_sec": round(time.time() - start, 2),
            })
    best = attempts[-1] if attempts else {}
    translations = best.get("translations", {})
    issues_by_lang = best.get("issues_by_lang") or {
        lang: validate_translation(en, translations.get(lang, ""), lang) or ["DIFY_ERROR"]
        for lang in langs
    }
    return {
        "en_hash": md5(en),
        "en": en,
        "langs": langs,
        "translations": translations,
        "issues_by_lang": issues_by_lang,
        "attempts": attempts,
        "status": "ok" if translations else "error",
        "ts": datetime.now(timezone.utc).isoformat(),
    }


def make_batches(pending, batch_size, max_batch_chars):
    grouped = defaultdict(list)
    for en, langs in pending:
        grouped[tuple(langs)].append({"en": en, "langs": langs})
    batches = []
    for langs, items in grouped.items():
        items.sort(key=lambda item: (len(item["en"]), item["en"]))
        current = []
        current_chars = 0
        for item in items:
            estimated = len(item["en"]) + 24
            if current and (len(current) >= batch_size or current_chars + estimated > max_batch_chars):
                batches.append((list(current), list(langs)))
                current = []
                current_chars = 0
            current.append(item)
            current_chars += estimated
        if current:
            batches.append((list(current), list(langs)))
    batches.sort(key=lambda item: (len(item[0][0]["en"]), len(item[0])))
    return batches


def translate_batch(batch, langs, api_key, primary_model, fallback_model, timeout):
    attempts = []
    for model in [primary_model, fallback_model]:
        if not model:
            continue
        start = time.time()
        try:
            translations_by_lang, raw = call_dify_batch(batch, langs, model, api_key, timeout)
            attempts.append({
                "model": model,
                "elapsed_sec": round(time.time() - start, 2),
                "raw": raw,
            })
            records = []
            all_good = True
            for idx, item in enumerate(batch):
                translations = {}
                issues_by_lang = {}
                for lang in langs:
                    translation = translations_by_lang.get(lang, {}).get(idx, "")
                    if lang == "sr":
                        translation = latinize_sr(translation)
                    translations[lang] = translation
                    issues = validate_translation(item["en"], translation, lang)
                    if not translation and "EMPTY" in issues:
                        issues.append("BATCH_PARSE_MISSING")
                    issues_by_lang[lang] = issues
                    if issues:
                        all_good = False
                records.append({
                    "en_hash": md5(item["en"]),
                    "en": item["en"],
                    "langs": langs,
                    "translations": translations,
                    "issues_by_lang": issues_by_lang,
                    "attempts": attempts,
                    "status": "ok",
                    "ts": datetime.now(timezone.utc).isoformat(),
                    "mode": "batch",
                })
            if all_good:
                return records
            if model == fallback_model:
                return records
        except Exception as exc:
            attempts.append({
                "model": model,
                "error": str(exc),
                "elapsed_sec": round(time.time() - start, 2),
            })
    return [{
        "en_hash": md5(item["en"]),
        "en": item["en"],
        "langs": langs,
        "translations": {},
        "issues_by_lang": {lang: ["DIFY_ERROR"] for lang in langs},
        "attempts": attempts,
        "status": "error",
        "ts": datetime.now(timezone.utc).isoformat(),
        "mode": "batch",
    } for item in batch]


def translate_tasks(workers, timeout, primary_model, fallback_model, batch_size, max_batch_chars):
    api_key = os.environ.get("DIFY_API_KEY") or os.environ.get("DIFY_NB_API_KEY")
    if not api_key:
        raise RuntimeError("缺少环境变量 DIFY_API_KEY 或 DIFY_NB_API_KEY")

    tasks = load_tasks()
    unique = {}
    for lang, rows in tasks.items():
        for row in rows:
            en = row["en"]
            unique.setdefault(md5(en), {"en": en, "langs": set()})
            unique[md5(en)]["langs"].add(lang)

    jsonl_path = PROCESS_DIR / "nb_sr_english_residual_dify_20260716.jsonl"
    done = load_done(jsonl_path)
    pending = []
    for en_hash, item in unique.items():
        if record_complete(done.get(en_hash), sorted(item["langs"])):
            continue
        pending.append((item["en"], sorted(item["langs"])))
    pending.sort(key=lambda item: (len(item[0]), item[0]))
    log(f"唯一英文 {len(unique)}，已完成 {len(done)}，待 Dify {len(pending)}")

    counters = defaultdict(int)
    if pending:
        if batch_size <= 1:
            work_items = [(en, langs) for en, langs in pending]
            total_units = len(work_items)
            with ThreadPoolExecutor(max_workers=workers) as executor:
                futures = {
                    executor.submit(translate_one, en, langs, api_key, primary_model, fallback_model, timeout): (en, langs)
                    for en, langs in work_items
                }
                for idx, future in enumerate(as_completed(futures), 1):
                    records = [future.result()]
                    for record in records:
                        append_jsonl(jsonl_path, record)
                        if record.get("status") == "ok":
                            bad = sum(1 for issues in record.get("issues_by_lang", {}).values() if issues)
                            counters["ok_with_issues" if bad else "ok"] += 1
                        else:
                            counters["error"] += 1
                    if idx % 50 == 0 or idx == total_units:
                        log(f"  Dify {idx}/{total_units} counters={dict(counters)}")
        else:
            work_items = make_batches(pending, batch_size, max_batch_chars)
            total_records = sum(len(batch) for batch, _ in work_items)
            completed_records = 0
            log(f"批量模式 batch_count={len(work_items)} batch_size={batch_size} total_records={total_records}")
            with ThreadPoolExecutor(max_workers=workers) as executor:
                futures = {
                    executor.submit(translate_batch, batch, langs, api_key, primary_model, fallback_model, timeout): (batch, langs)
                    for batch, langs in work_items
                }
                for idx, future in enumerate(as_completed(futures), 1):
                    records = future.result()
                    for record in records:
                        append_jsonl(jsonl_path, record)
                        completed_records += 1
                        if record.get("status") == "ok":
                            bad = sum(1 for issues in record.get("issues_by_lang", {}).values() if issues)
                            counters["ok_with_issues" if bad else "ok"] += 1
                        else:
                            counters["error"] += 1
                    if idx % 10 == 0 or idx == len(work_items):
                        log(f"  Dify batch {idx}/{len(work_items)} records={completed_records}/{total_records} counters={dict(counters)}")
    return str(jsonl_path)


def set_text(cell, value):
    text = clean(value)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"
    if text.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def write_upload_file(path, lang, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = lang
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165
    for col_idx, header in enumerate(["key ID", "en", lang, "备注"], 1):
        cell = ws.cell(2, col_idx)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 3):
        values = [row["keyId"], row["en"], row[lang], ""]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 30, "B": 72, "C": 72, "D": 18}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def split_write_uploads(lang, rows):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    if not rows:
        return paths
    for idx in range(0, len(rows), MAX_UPLOAD_ROWS):
        part_rows = rows[idx:idx + MAX_UPLOAD_ROWS]
        part_no = idx // MAX_UPLOAD_ROWS + 1
        suffix = f"_part{part_no:02d}" if len(rows) > MAX_UPLOAD_ROWS else ""
        path = OUTPUT_DIR / f"{lang}_英文残留重翻_可上传_20260716{suffix}.xlsx"
        write_upload_file(path, lang, part_rows)
        paths.append(str(path))
    return paths


def build_outputs():
    tasks = load_tasks()
    jsonl_path = PROCESS_DIR / "nb_sr_english_residual_dify_20260716.jsonl"
    done = load_done(jsonl_path)

    good = {lang: [] for lang in LANGS}
    unresolved = {lang: [] for lang in LANGS}
    seen = {lang: set() for lang in LANGS}

    for lang, rows in tasks.items():
        for row in rows:
            en_hash = md5(row["en"])
            record = done.get(en_hash)
            translation = ""
            issues = ["NO_DIFY_RESULT"]
            if record:
                translation = norm((record.get("translations") or {}).get(lang))
                if lang == "sr":
                    translation = latinize_sr(translation)
                issues = validate_translation(row["en"], translation, lang)
            if issues:
                bad = dict(row)
                bad[lang] = translation
                bad["issues"] = issues
                unresolved[lang].append(bad)
                continue
            dedupe_key = (row["keyId"], row["en"])
            if dedupe_key in seen[lang]:
                continue
            seen[lang].add(dedupe_key)
            good[lang].append({
                "keyId": row["keyId"],
                "en": row["en"],
                lang: translation,
                "source_file": row["source_file"],
                "reason": row["reason"],
            })

    upload_paths = {lang: split_write_uploads(lang, good[lang]) for lang in LANGS}
    unresolved_path = OUTPUT_DIR / "nb_sr_英文残留重翻_仍需处理_20260716.xlsx"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    first = True
    for lang in LANGS:
        ws = wb.active if first else wb.create_sheet(lang)
        first = False
        ws.title = lang
        headers = ["key ID", "en", lang, "问题类型", "reason", "source_file", "sheet", "row"]
        ws.append(headers)
        for col_idx, _ in enumerate(headers, 1):
            cell = ws.cell(1, col_idx)
            cell.font = Font(bold=True)
            cell.fill = BAD_FILL
        for item in unresolved[lang]:
            ws.append([
                item["keyId"],
                item["en"],
                item.get(lang, ""),
                ";".join(item.get("issues", [])),
                item.get("reason", ""),
                item.get("source_file", ""),
                item.get("sheet", ""),
                item.get("row", ""),
            ])
        for col, width in {"A": 30, "B": 72, "C": 72, "D": 30, "E": 24, "F": 44, "G": 18, "H": 10}.items():
            ws.column_dimensions[col].width = width
    wb.save(unresolved_path)

    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "upload_paths": upload_paths,
        "good_rows": {lang: len(good[lang]) for lang in LANGS},
        "unresolved_rows": {lang: len(unresolved[lang]) for lang in LANGS},
        "unresolved_path": str(unresolved_path),
    }
    report_path = PROCESS_DIR / "nb_sr_english_residual_output_report_20260716.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    log(json.dumps(report, ensure_ascii=False, indent=2))
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--timestamp", help="全量导出时间戳，如 20260716_151753")
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--translate-only", action="store_true")
    parser.add_argument("--build-only", action="store_true")
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--timeout", type=int, default=180)
    parser.add_argument("--primary-model", default=DEFAULT_MODEL)
    parser.add_argument("--fallback-model", default=FALLBACK_MODEL)
    parser.add_argument("--batch-size", type=int, default=12)
    parser.add_argument("--max-batch-chars", type=int, default=2500)
    args = parser.parse_args()

    if args.build_only:
        build_outputs()
        return
    if not args.translate_only:
        summary = scan_tasks(args.timestamp)
        log(json.dumps(summary, ensure_ascii=False, indent=2))
    if args.prepare_only:
        return
    translate_tasks(args.workers, args.timeout, args.primary_model, args.fallback_model, args.batch_size, args.max_batch_chars)
    build_outputs()


if __name__ == "__main__":
    main()
