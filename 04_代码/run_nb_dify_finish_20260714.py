#!/usr/bin/env python3
"""用富文本保护 Dify 补齐 nb 剩余内容，并汇总最终上传文件。

凭据通过环境变量 DIFY_NB_API_KEY 传入，不写入脚本。
"""

import argparse
import hashlib
import json
import os
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
TASK_JSON = BASE / "02_过程数据/refill_tasks_by_lang_20260713/nb_refill_tasks_20260713.json"
CANDIDATE_XLSX = BASE / "05_上传格式/20260713_补翻/nb_自动修复后_可上传候选_20260714.xlsx"
STILL_BAD_XLSX = BASE / "05_上传格式/20260713_补翻/nb_自动修复后_仍需处理_20260714.xlsx"

PROCESS = BASE / "02_过程数据"
JSONL = PROCESS / "nb_dify_finish_translations_20260714.jsonl"
REPORT_JSON = PROCESS / "nb_dify_finish_report_20260714.json"
FINAL_XLSX = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_最终可上传_20260714.xlsx"
UNRESOLVED_XLSX = BASE / "05_上传格式/20260713_补翻/nb_Dify补齐_仍需处理_20260714.xlsx"

DIFY_PROXY = os.environ.get("DIFY_PROXY") or "http://127.0.0.1:30001/rest"
DIFY_PATH = "/aihelper-dify-api2/v1/workflows/run"
TEXT_CONTEXT = (
    "SHEIN e-commerce UI copy. Target language code nb means Norwegian Bokmål. "
    "Preserve placeholders such as {0}, %s and all HTML tags exactly, including order and attributes. "
    "Do not translate placeholder names inside braces."
)

PH_RE = re.compile(r"(\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+)")
TAG_RE = re.compile(r"<[^>]+>")
FORMULA_PREFIXES = ("=", "+", "-", "@")
write_lock = threading.Lock()

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)


def log(message):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def norm(value):
    return "" if value is None else str(value)


def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub("", norm(value))


def md5(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def placeholders(text):
    return PH_RE.findall(norm(text))


def tags(text):
    return TAG_RE.findall(norm(text))


def validate(en, nb):
    issues = []
    if not nb.strip():
        issues.append("EMPTY_NB")
    if sorted(placeholders(en)) != sorted(placeholders(nb)):
        issues.append("PH_MISMATCH")
    if tags(en) and tags(en) != tags(nb):
        issues.append("HTML_TAG_MISMATCH")
    return issues


def set_text(cell, value):
    text = clean(value)
    cell.value = text
    cell.data_type = "s"
    cell.number_format = "@"
    if text.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def rows_from_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return list(ws.iter_rows(values_only=True))


def load_task_rows():
    return json.loads(TASK_JSON.read_text(encoding="utf-8"))


def load_candidate_rows():
    rows = rows_from_xlsx(CANDIDATE_XLSX)
    headers = [norm(value).strip() for value in rows[1]]
    key_idx = headers.index("key ID")
    en_idx = headers.index("en")
    nb_idx = headers.index("nb")
    result = {}
    for row in rows[2:]:
        key = norm(row[key_idx]).strip()
        if not key:
            continue
        result[key] = {
            "keyId": key,
            "en": norm(row[en_idx]).strip(),
            "nb": norm(row[nb_idx]).strip(),
            "source": "candidate",
        }
    return result


def load_still_bad_rows():
    rows = rows_from_xlsx(STILL_BAD_XLSX)
    headers = [norm(value).strip() for value in rows[0]]
    key_idx = headers.index("key ID")
    en_idx = headers.index("en")
    old_idx = headers.index("自动修复后nb")
    issue_idx = headers.index("修复后问题")
    result = {}
    for row in rows[1:]:
        key = norm(row[key_idx]).strip()
        if not key:
            continue
        result[key] = {
            "keyId": key,
            "en": norm(row[en_idx]).strip(),
            "old_nb": norm(row[old_idx]).strip(),
            "old_issue": norm(row[issue_idx]).strip(),
            "source": "still_bad",
        }
    return result


def load_done_jsonl():
    done = {}
    if not JSONL.exists():
        return done
    with open(JSONL, encoding="utf-8") as f:
        for line in f:
            try:
                record = json.loads(line)
            except Exception:
                continue
            if record.get("status") != "ok":
                continue
            key = record.get("en_hash")
            if key and record.get("nb"):
                done[key] = record
    return done


def append_jsonl(record):
    with write_lock:
        with open(JSONL, "a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def extract_translation(resp):
    outputs = (resp.get("data") or {}).get("outputs") or {}
    output = outputs.get("output")
    if isinstance(output, list) and output:
        return norm(output[0].get("translated_text")).strip(), output[0]
    if isinstance(output, str):
        return output.strip(), {"output": output}
    return "", outputs


def call_dify(en, model, api_key, timeout):
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "mcp-proxy-path": DIFY_PATH,
        "mcp-proxy-env": "prod",
        "mcp-proxy-dc": "aliyun_dc",
    }
    payload = {
        "inputs": {
            "model": model,
            "OriginalLanguage": "English",
            "OriginalText": en,
            "TargetLanguage": "nb",
            "TextContext": TEXT_CONTEXT,
        },
        "response_mode": "blocking",
        "user": "codex-nb-dify-finish",
    }
    response = requests.post(
        DIFY_PROXY,
        headers=headers,
        data=json.dumps(payload, ensure_ascii=False),
        timeout=timeout,
    )
    response.raise_for_status()
    text, raw = extract_translation(response.json())
    return text, raw


def translate_one(en, api_key, primary_model, fallback_model, timeout):
    attempts = []
    for model in [primary_model, fallback_model]:
        if not model:
            continue
        start = time.time()
        try:
            nb, raw = call_dify(en, model, api_key, timeout)
            issues = validate(en, nb)
            attempts.append({
                "model": model,
                "nb": nb,
                "issues": issues,
                "elapsed_sec": round(time.time() - start, 2),
                "raw": raw,
            })
            if not issues:
                return {
                    "en_hash": md5(en),
                    "en": en,
                    "nb": nb,
                    "issues": [],
                    "attempts": attempts,
                    "status": "ok",
                    "ts": datetime.now(timezone.utc).isoformat(),
                }
        except Exception as exc:
            attempts.append({
                "model": model,
                "error": str(exc),
                "elapsed_sec": round(time.time() - start, 2),
            })
    best = attempts[-1] if attempts else {}
    return {
        "en_hash": md5(en),
        "en": en,
        "nb": best.get("nb", ""),
        "issues": best.get("issues", ["DIFY_ERROR"]),
        "attempts": attempts,
        "status": "ok" if best.get("nb") else "error",
        "ts": datetime.now(timezone.utc).isoformat(),
    }


def write_upload_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "nb"
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165
    for col, header in enumerate(["key ID", "en", "nb", "备注"], 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 3):
        for col_idx, value in enumerate([row["keyId"], row["en"], row["nb"], ""], 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 28, "B": 70, "C": 70, "D": 20}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def write_unresolved_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "仍需处理"
    headers = ["key ID", "en", "nb", "问题类型", "来源"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col)
        set_text(cell, header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F4CCCC")
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(rows, 2):
        values = [
            row["keyId"], row["en"], row.get("nb", ""),
            ";".join(row.get("issues", [])), row.get("source", ""),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 28, "B": 70, "C": 70, "D": 28, "E": 18}.items():
        ws.column_dimensions[col].width = width
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--timeout", type=int, default=180)
    parser.add_argument("--primary-model", default="gemini-3-flash")
    parser.add_argument("--fallback-model", default="gpt-5-mini")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--skip-translate", action="store_true")
    args = parser.parse_args()

    api_key = os.environ.get("DIFY_NB_API_KEY")
    if not api_key:
        raise RuntimeError("缺少环境变量 DIFY_NB_API_KEY")

    task_rows = load_task_rows()
    candidate_by_key = load_candidate_rows()
    still_bad_by_key = load_still_bad_rows()
    missing_keys = [
        row["keyId"]
        for row in task_rows
        if row["keyId"] not in candidate_by_key and row["keyId"] not in still_bad_by_key
    ]
    key_to_task = {row["keyId"]: row for row in task_rows}

    translate_items = {}
    for key, row in still_bad_by_key.items():
        translate_items[md5(row["en"])] = row["en"]
    for key in missing_keys:
        row = key_to_task[key]
        translate_items[md5(row["en"])] = row["en"]

    done = load_done_jsonl()
    pending = [(h, en) for h, en in translate_items.items() if h not in done]
    pending.sort(key=lambda item: (len(item[1]), item[1]))
    if args.limit:
        pending = pending[:args.limit]
    log(
        f"candidate={len(candidate_by_key)} still_bad={len(still_bad_by_key)} "
        f"missing_empty={len(missing_keys)} unique_translate={len(translate_items)} "
        f"pending={len(pending)}"
    )

    if pending and not args.skip_translate:
        counters = {"ok": 0, "error": 0}

        def work(item):
            _, en = item
            return translate_one(en, api_key, args.primary_model, args.fallback_model, args.timeout)

        with ThreadPoolExecutor(max_workers=args.workers) as executor:
            futures = {executor.submit(work, item): item for item in pending}
            for idx, future in enumerate(as_completed(futures), 1):
                en_hash, en = futures[future]
                try:
                    record = future.result()
                except Exception as exc:
                    record = {
                        "en_hash": en_hash,
                        "en": en,
                        "nb": "",
                        "issues": ["DIFY_EXCEPTION"],
                        "status": "error",
                        "error": str(exc),
                        "ts": datetime.now(timezone.utc).isoformat(),
                    }
                append_jsonl(record)
                done[record["en_hash"]] = record
                if record.get("status") == "ok" and not record.get("issues"):
                    counters["ok"] += 1
                else:
                    counters["error"] += 1
                if idx <= 5 or idx % 20 == 0:
                    log(f"进度 {idx}/{len(pending)} ok={counters['ok']} error={counters['error']}")

    done = load_done_jsonl()
    final_rows = []
    unresolved = []
    for source_row in task_rows:
        key = source_row["keyId"]
        en = source_row["en"]
        if key in candidate_by_key:
            row = candidate_by_key[key]
            final_rows.append({"keyId": key, "en": row["en"], "nb": row["nb"], "source": "candidate"})
            continue
        record = done.get(md5(en), {})
        nb = record.get("nb", "")
        issues = validate(en, nb)
        if nb and not issues:
            final_rows.append({"keyId": key, "en": en, "nb": nb, "source": "dify"})
        else:
            unresolved.append({
                "keyId": key,
                "en": en,
                "nb": nb,
                "issues": issues or record.get("issues") or ["NO_TRANSLATION"],
                "source": "dify",
            })

    write_upload_workbook(FINAL_XLSX, final_rows)
    write_unresolved_workbook(UNRESOLVED_XLSX, unresolved)

    reason_counts = {}
    for row in unresolved:
        reason = ";".join(row["issues"])
        reason_counts[reason] = reason_counts.get(reason, 0) + 1

    report = {
        "task_rows": len(task_rows),
        "candidate_rows": len(candidate_by_key),
        "still_bad_rows": len(still_bad_by_key),
        "missing_empty_rows": len(missing_keys),
        "unique_translate": len(translate_items),
        "final_upload_rows": len(final_rows),
        "unresolved_rows": len(unresolved),
        "unresolved_reason_counts": reason_counts,
        "jsonl": str(JSONL),
        "final_upload_file": str(FINAL_XLSX),
        "unresolved_file": str(UNRESOLVED_XLSX),
    }
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
