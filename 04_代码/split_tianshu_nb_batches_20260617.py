#!/usr/bin/env python3
"""
把 nb 待翻译唯一英文拆成天枢批量任务输入文件。

输入模板：/Users/10304327/Downloads/input.xlsx，仅使用其 text 表头。
数据来源：02_过程数据/nb_full_20260616_bokmal_dify_20260617_tasks.json
输出：02_过程数据/tianshu_nb_batches_20260617/*.xlsx
"""

import json
import os
from pathlib import Path

import openpyxl
from openpyxl import Workbook

BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
TEMPLATE = Path("/Users/10304327/Downloads/input.xlsx")
TASK_JSON = BASE / "02_过程数据" / "nb_full_20260616_bokmal_dify_20260617_tasks.json"
OUT_DIR = BASE / "02_过程数据" / "tianshu_nb_batches_20260617"
MANIFEST = OUT_DIR / "manifest.json"

MAX_BYTES = 4_900_000
MAX_ROWS = 50_000
TARGET_TEXT_BYTES = 4_000_000
HEADER = "text"


def read_template_header():
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True, data_only=True)
    ws = wb.active
    value = ws.cell(row=1, column=1).value
    header = "" if value is None else str(value).strip()
    if header != HEADER:
        raise ValueError(f"模板首列表头不是 {HEADER!r}: {header!r}")
    return header


def load_texts():
    data = json.loads(TASK_JSON.read_text(encoding="utf-8"))
    texts = [str(text) for text in data["needed"] if str(text).strip()]
    seen = set()
    deduped = []
    for text in texts:
        if text not in seen:
            seen.add(text)
            deduped.append(text)
    return deduped


def write_xlsx(path, texts):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append([HEADER])
    for text in texts:
        ws.append([text])
    wb.save(path)
    return path.stat().st_size


def split_by_text_bytes(texts):
    batches = []
    current = []
    current_bytes = 0
    for text in texts:
        size = len(text.encode("utf-8")) + 16
        if current and (current_bytes + size > TARGET_TEXT_BYTES or len(current) >= MAX_ROWS):
            batches.append(current)
            current = []
            current_bytes = 0
        current.append(text)
        current_bytes += size
    if current:
        batches.append(current)
    return batches


def ensure_under_limit(path, texts, batch_no):
    size = write_xlsx(path, texts)
    if size <= MAX_BYTES or len(texts) == 1:
        return [{"path": path, "rows": len(texts), "bytes": size}]

    path.unlink()
    mid = len(texts) // 2
    left = OUT_DIR / f"nb_tianshu_{batch_no:03d}a.xlsx"
    right = OUT_DIR / f"nb_tianshu_{batch_no:03d}b.xlsx"
    return ensure_under_limit(left, texts[:mid], batch_no) + ensure_under_limit(right, texts[mid:], batch_no)


def verify_xlsx(path, expected_rows):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    row_iter = ws.iter_rows(values_only=True)
    first_row = next(row_iter, None)
    header = None if not first_row else first_row[0]
    rows = sum(1 for _ in row_iter)
    if header != HEADER:
        raise ValueError(f"{path.name} 表头异常: {header!r}")
    if rows != expected_rows:
        raise ValueError(f"{path.name} 行数异常: expected={expected_rows}, actual={rows}")
    if rows > MAX_ROWS:
        raise ValueError(f"{path.name} 超过 50000 行限制: {rows}")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError(f"{path.name} 超过 5MB 限制: {path.stat().st_size}")


def main():
    read_template_header()
    texts = load_texts()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    manifest = {
        "template": str(TEMPLATE),
        "task_json": str(TASK_JSON),
        "target_lang": "nb",
        "target_lang_name": "挪威布克莫尔语（书面挪威语）",
        "max_bytes": MAX_BYTES,
        "max_rows": MAX_ROWS,
        "total_unique_texts": len(texts),
        "files": [],
    }

    for idx, batch in enumerate(split_by_text_bytes(texts), 1):
        path = OUT_DIR / f"nb_tianshu_{idx:03d}.xlsx"
        manifest["files"].extend(ensure_under_limit(path, batch, idx))

    total_rows = 0
    for item in manifest["files"]:
        verify_xlsx(item["path"], item["rows"])
        total_rows += item["rows"]
        item["path"] = str(item["path"])

    if total_rows != len(texts):
        raise ValueError(f"总行数异常: expected={len(texts)}, actual={total_rows}")

    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
