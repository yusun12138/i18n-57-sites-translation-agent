#!/usr/bin/env python3
"""按今天最新全量对账最终待上传文件，生成需要重跑的明细。"""

import csv
import json
import os
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl


BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
RAW = BASE / "01_原始数据"
UPLOAD = BASE / "05_上传格式"
PROCESS = BASE / "02_过程数据"
TASK_DIR = PROCESS / "retry_tasks_by_lang_20260624"
EXCLUDED = BASE / "04_代码" / "excluded_keyids.txt"
RUN_DATE = "20260624"

# 用户确认：zh-hk/en_au/en_gb/zh 不涉及；en 是源文列，不作为目标语种。
TARGET_LANGS = [
    "nb",
    "ru",
    "ar",
    "lt",
    "id",
    "el-gr",
    "nl",
    "sl",
    "it",
    "hu",
    "tr",
    "pt-pt",
    "lv",
    "uk",
    "pl",
    "ro",
    "es",
    "fr",
    "sr",
    "de",
    "et",
    "zh-cn",
    "zh-tw",
]

FORMULA_PREFIXES = ("=", "+", "-", "@")
META_COLS = {
    "id",
    "keyId",
    "en",
    "所属产品",
    "备注",
    "产品创建人",
    "key创建人",
    "key最后修改人",
}


def norm(value):
    return "" if value is None else str(value).strip()


def is_empty(value):
    return norm(value).lower() in {"", "none", "null"}


def load_excluded():
    if not EXCLUDED.exists():
        return set()
    return {
        line.strip()
        for line in EXCLUDED.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    }


def latest(pattern):
    files = sorted(RAW.glob(pattern), key=lambda path: path.stat().st_mtime)
    if not files:
        raise FileNotFoundError(pattern)
    return files[-1]


def latest_input_files():
    files = [
        latest(f"type1_前端key_platform{platform}_full_{RUN_DATE}_*.xlsx")
        for platform in [1, 2, 3, 4, 5]
    ]
    files.append(latest(f"type2_提示语_full_{RUN_DATE}_*.xlsx"))
    type3_files = sorted(
        list(RAW.glob(f"type3_错误码_full_{RUN_DATE}_*.xlsx"))
        + list(RAW.glob(f"type3_错误码_{RUN_DATE}_*.xlsx")),
        key=lambda path: path.stat().st_mtime,
    )
    if not type3_files:
        raise FileNotFoundError(f"type3_错误码*{RUN_DATE}*.xlsx")
    files.append(type3_files[-1])
    return files


def load_online_missing():
    """返回 lang -> keyId -> latest row info，只保留线上目标语种为空的行。"""
    excluded = load_excluded()
    missing = {lang: {} for lang in TARGET_LANGS}
    source_files = []
    total_rows = 0
    rows_with_en = 0

    for path in latest_input_files():
        source_files.append(str(path))
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()
        iterator = ws.iter_rows(values_only=True)
        headers = [norm(value) for value in next(iterator)]
        lower_headers = [header.lower() for header in headers]
        key_idx = next(i for i, header in enumerate(headers) if header in {"keyId", "key ID", "key_id"})
        en_idx = lower_headers.index("en")
        lang_idx = {lang: headers.index(lang) for lang in TARGET_LANGS if lang in headers}

        for row in iterator:
            total_rows += 1
            key = norm(row[key_idx] if key_idx < len(row) else "")
            en = norm(row[en_idx] if en_idx < len(row) else "")
            if not key or not en:
                continue
            rows_with_en += 1
            if key in excluded:
                continue
            for lang, idx in lang_idx.items():
                value = row[idx] if idx < len(row) else None
                if is_empty(value):
                    missing[lang][key] = {
                        "keyId": key,
                        "latest_en": en,
                        "source_file": path.name,
                    }

    return {
        "source_files": source_files,
        "total_rows": total_rows,
        "rows_with_en": rows_with_en,
        "missing": missing,
    }


def upload_candidates():
    files = []
    for path in sorted(UPLOAD.glob("*.xlsx")):
        if path.name.startswith(".~") or "有问题" in path.name:
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            next(rows, None)
            header = next(rows, None)
            if not header:
                continue
            headers = [norm(value) for value in header]
            langs = [lang for lang in TARGET_LANGS if lang in headers]
            if langs:
                files.append({"path": path, "headers": headers, "langs": langs})
        except Exception as exc:
            files.append({"path": path, "headers": [], "langs": [], "error": str(exc)})
    return files


def load_upload_records():
    """返回 lang -> keyId -> 上传记录；同 key 多文件时保留“最好”的记录。"""
    records = {lang: {} for lang in TARGET_LANGS}
    files = upload_candidates()
    file_stats = []

    def quality(record):
        if record["target"] and not record["remark"]:
            return 3
        if record["target"] and record["remark"]:
            return 2
        if not record["target"]:
            return 1
        return 0

    for item in files:
        path = item["path"]
        if item.get("error") or not item["langs"]:
            continue
        wb = openpyxl.load_workbook(path, data_only=False)
        ws = wb.active
        headers = [norm(ws.cell(2, col).value) for col in range(1, ws.max_column + 1)]
        remark_col = headers.index("备注") + 1 if "备注" in headers else ws.max_column
        stat = {
            "file": path.name,
            "langs": item["langs"],
            "rows": 0,
            "target_empty_rows": 0,
            "remark_rows": 0,
            "formula_cell_count": 0,
            "formula_prefix_without_quote_count": 0,
        }

        for row_idx in range(3, ws.max_row + 1):
            key = norm(ws.cell(row_idx, 1).value)
            upload_en = norm(ws.cell(row_idx, 2).value)
            if not key:
                continue
            stat["rows"] += 1
            remark = norm(ws.cell(row_idx, remark_col).value)
            if remark:
                stat["remark_rows"] += 1

            for lang in item["langs"]:
                lang_col = headers.index(lang) + 1
                cell = ws.cell(row_idx, lang_col)
                target = norm(cell.value)
                if not target:
                    stat["target_empty_rows"] += 1
                if cell.data_type == "f":
                    stat["formula_cell_count"] += 1
                if isinstance(cell.value, str) and cell.value.startswith(FORMULA_PREFIXES) and not cell.quotePrefix:
                    stat["formula_prefix_without_quote_count"] += 1

                record = {
                    "keyId": key,
                    "upload_en": upload_en,
                    "target": target,
                    "remark": remark,
                    "upload_file": path.name,
                    "upload_row": row_idx,
                }
                existing = records[lang].get(key)
                if existing is None or quality(record) > quality(existing):
                    records[lang][key] = record

        file_stats.append(stat)

    return {
        "records": records,
        "files": [
            {
                "file": item["path"].name,
                "langs": item["langs"],
                **({"error": item["error"]} if item.get("error") else {}),
            }
            for item in files
        ],
        "file_stats": file_stats,
    }


def classify(online, upload):
    if upload is None:
        return "NOT_IN_FINAL_UPLOAD"
    if not upload["target"]:
        return "EMPTY_TARGET_IN_UPLOAD"
    if upload["remark"]:
        return "RISK_REMARK_IN_UPLOAD"
    if upload["upload_en"] != online["latest_en"]:
        return "EN_CHANGED"
    return "READY_TO_UPLOAD"


def main():
    online_data = load_online_missing()
    upload_data = load_upload_records()
    upload_records = upload_data["records"]

    retry_rows = []
    ready_rows = []
    summary = {}

    for lang in TARGET_LANGS:
        reason_counts = Counter()
        missing_rows = online_data["missing"][lang]
        for key, online in sorted(missing_rows.items()):
            upload = upload_records[lang].get(key)
            reason = classify(online, upload)
            reason_counts[reason] += 1
            row = {
                "lang": lang,
                "keyId": key,
                "latest_en": online["latest_en"],
                "upload_en": upload["upload_en"] if upload else "",
                "old_translation": upload["target"] if upload else "",
                "retry_reason": reason,
                "source_file": online["source_file"],
                "upload_file": upload["upload_file"] if upload else "",
                "upload_row": upload["upload_row"] if upload else "",
                "remark": upload["remark"] if upload else "",
            }
            if reason == "READY_TO_UPLOAD":
                ready_rows.append(row)
            else:
                retry_rows.append(row)

        summary[lang] = {
            "online_missing": len(missing_rows),
            "upload_records": len(upload_records[lang]),
            "reason_counts": dict(reason_counts),
            "retry_total": sum(count for reason, count in reason_counts.items() if reason != "READY_TO_UPLOAD"),
            "ready_to_upload": reason_counts.get("READY_TO_UPLOAD", 0),
        }

    PROCESS.mkdir(parents=True, exist_ok=True)
    TASK_DIR.mkdir(parents=True, exist_ok=True)

    retry_csv = PROCESS / f"retry_tasks_{RUN_DATE}.csv"
    ready_csv = PROCESS / f"ready_to_upload_coverage_{RUN_DATE}.csv"
    fieldnames = [
        "lang",
        "keyId",
        "latest_en",
        "upload_en",
        "old_translation",
        "retry_reason",
        "source_file",
        "upload_file",
        "upload_row",
        "remark",
    ]
    for path, rows in [(retry_csv, retry_rows), (ready_csv, ready_rows)]:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    by_lang = defaultdict(list)
    for row in retry_rows:
        by_lang[row["lang"]].append(
            {
                "keyId": row["keyId"],
                "en": row["latest_en"],
                "retry_reason": row["retry_reason"],
                "source_file": row["source_file"],
                "old_translation": row["old_translation"],
            }
        )
    task_files = {}
    for lang, rows in by_lang.items():
        path = TASK_DIR / f"{lang}_retry_tasks_{RUN_DATE}.json"
        path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
        task_files[lang] = str(path)

    report = {
        "run_date": RUN_DATE,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "target_langs": TARGET_LANGS,
        "source_files": online_data["source_files"],
        "upload_files": upload_data["files"],
        "upload_file_stats": upload_data["file_stats"],
        "summary": summary,
        "retry_csv": str(retry_csv),
        "ready_csv": str(ready_csv),
        "task_files": task_files,
    }
    report_path = PROCESS / f"upload_precheck_report_{RUN_DATE}.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps({
        "report": str(report_path),
        "retry_csv": str(retry_csv),
        "ready_csv": str(ready_csv),
        "summary": summary,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
