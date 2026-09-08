#!/usr/bin/env python3
"""
独立 nb 书面挪威语重翻脚本。

不改动 20260616 现有翻译脚本和旧 JSONL。调用接口前需要设置 DIFY_API_KEY。
默认只探测接口：
  python3 04_代码/translate_nb_bokmal_dify_20260617.py --action probe

小样本验证：
  python3 04_代码/translate_nb_bokmal_dify_20260617.py --action translate --limit 10 --workers 1

生成上传文件：
  python3 04_代码/translate_nb_bokmal_dify_20260617.py --action build-upload
"""

import argparse
import hashlib
import json
import os
import random
import re
import threading
import time
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(os.environ.get("I18N_WORKDIR") or Path(__file__).resolve().parents[1]).resolve()
RAW = BASE / "01_原始数据"
PROCESS = BASE / "02_过程数据"
UPLOAD = BASE / "05_上传格式"
EXCLUDED = BASE / "04_代码" / "excluded_keyids.txt"

SOURCE_DATE = os.environ.get("SOURCE_DATE") or "20260616"
RUN_TAG = os.environ.get("NB_RUN_TAG") or "bokmal_dify_20260617"
LANG = "nb"
TARGET_CODE = "nb"
TARGET_NAME = "挪威布克莫尔语（书面挪威语）"

DIFY_PROXY = os.environ.get("DIFY_PROXY") or "http://127.0.0.1:30001/rest"
DIFY_HEALTH = os.environ.get("DIFY_HEALTH") or "http://127.0.0.1:30001/health"
DIFY_API_KEY = os.environ.get("DIFY_API_KEY", "").strip()
DIFY_HEADERS_BASE = {
    "Content-Type": "application/json",
    "mcp-proxy-path": "/aihelper-dify-api2/v1/workflows/run",
    "mcp-proxy-env": "prod",
    "mcp-proxy-dc": "aliyun_dc",
}

JSONL = PROCESS / f"nb_full_{SOURCE_DATE}_{RUN_TAG}.jsonl"
TASK_CACHE = PROCESS / f"nb_full_{SOURCE_DATE}_{RUN_TAG}_tasks.json"
OUT_XLSX = UPLOAD / f"nb_{RUN_TAG}.xlsx"
QC_JSON = PROCESS / f"nb_{RUN_TAG}_upload_qc.json"

PH_RE = re.compile(r"\{[^}]+\}|%\d+\$[sdf]|%[sdf]|%\d+")
TAG_RE = re.compile(r"<[^>]+>")
ESC_RE = re.compile(r"\\['\u2018\u2019]")
CURLY_RE = re.compile(r"\{([^}]*)\}")
FORMULA_PREFIXES = ("=", "+", "-", "@")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
RISK_FILL = PatternFill("solid", fgColor="FFF2CC")
FLUSH_EVERY = 200

_buffer_lock = threading.Lock()
_pending_records = []

NOTICE = (
    "模版使用注意事项：\n"
    "1. 模版注意事项内容无需删掉\n"
    "2. 必填字段：key ID（key ID/code ID）；en（系统中原英文文案）；更新语种（需要更新的语种文案，语种名称需和系统中保持一致）\n"
    "3. 字段解释：\n"
    "key ID——需要更新的key ID或code ID，需和系统中的ID保持一致，上传时会进行校验\n"
    "en——需要更新的key在系统中的原英文文案\n"
    "更新语种——需要更新的语种名称，将字段替换为系统中的语种名称，上传时会进行校验\n"
    "4. 表格中仅填写，ID、原英文文案及需要更新的语种文案，其余语种文案不要填写，更新时会覆盖！\n"
    "例如，需要更新日语文案，则表格内容共三列：key ID、en、ja\n"
    "如果需要更新英语、韩语、阿语文案，则表格内容共五列：key ID、en、en、ko、ar（第一列en为原英文文案，第二列en为需要更新的en文案）\n"
    "5. 备注字段如果为空，不会覆盖线上数据。如果填写值，会覆盖线上数据。"
)


def log(message):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def norm(value):
    return "" if value is None else str(value).strip()


def md5(text):
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def escape_curly(text):
    return CURLY_RE.sub(lambda m: f"[[CURLY_{m.group(1)}]]", text)


def unescape_curly(text):
    return re.sub(r"\[\[CURLY_([^\]]*)\]\]", lambda m: "{" + m.group(1) + "}", text)


def protect(text):
    mapping = []

    def repl(match):
        token = f"ZXSHEIN{len(mapping)}ZX"
        mapping.append((token, match.group(0)))
        return token

    text = TAG_RE.sub(repl, text)
    text = PH_RE.sub(repl, text)
    return text, mapping


def restore(text, mapping):
    for token, raw in mapping:
        text = text.replace(token, raw)
    return text


def qc_issues(en, tr):
    en = str(en or "")
    tr = str(tr or "")
    issues = []
    if not tr.strip():
        issues.append("EMPTY")
    if sorted(PH_RE.findall(en)) != sorted(PH_RE.findall(tr)):
        issues.append("PH_MISMATCH")
    if ESC_RE.findall(en) != ESC_RE.findall(tr):
        issues.append("ESCAPED_SINGLE_QUOTE_MISMATCH")
    if TAG_RE.findall(en) != TAG_RE.findall(tr):
        issues.append("TAG_SEQUENCE_MISMATCH")
    if len(en) >= 1200 and len(tr) < len(en) * 0.45:
        issues.append(f"POSSIBLE_TRUNCATED_LONG_TEXT en_len={len(en)} got_len={len(tr)}")
    return issues


def qc_ok(en, tr):
    return not qc_issues(en, tr)


def should_copy_source(en):
    return str(en or "").strip() == "'"


def source_files():
    files = []
    for platform in [1, 2, 3, 4, 5]:
        matches = sorted(
            RAW.glob(f"type1_前端key_platform{platform}_full_{SOURCE_DATE}_*.xlsx"),
            key=lambda p: p.stat().st_mtime,
        )
        if not matches:
            raise FileNotFoundError(f"未找到 type1 platform{platform} 源文件，日期: {SOURCE_DATE}")
        files.append(matches[-1])
    for pattern in [
        f"type2_提示语_full_{SOURCE_DATE}_*.xlsx",
        f"type3_错误码_full_{SOURCE_DATE}_*.xlsx",
    ]:
        matches = sorted(RAW.glob(pattern), key=lambda p: p.stat().st_mtime)
        if not matches:
            raise FileNotFoundError(f"未找到源文件: {pattern}")
        files.append(matches[-1])
    return files


def load_excluded():
    if not EXCLUDED.exists():
        return set()
    return {
        line.strip()
        for line in EXCLUDED.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.strip().startswith("#")
    }


def find_lang_col(headers, lang):
    matches = [idx for idx, header in enumerate(headers) if header == lang]
    if lang == "id" and matches:
        return matches[-1]
    return matches[0] if matches else None


def build_needed_and_rows(force_refresh=False):
    if TASK_CACHE.exists() and not force_refresh:
        with open(TASK_CACHE, encoding="utf-8") as f:
            data = json.load(f)
        return set(data["needed"]), OrderedDict(data["missing_rows"]), data["source_stats"]

    excluded = load_excluded()
    needed = set()
    missing_rows = OrderedDict()
    source_stats = []

    for path in source_files():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = [norm(v).lower() for v in next(rows)]
        key_i = next(i for i, h in enumerate(headers) if h in ("keyid", "key id", "key_id"))
        en_i = headers.index("en")
        lang_i = find_lang_col(headers, LANG)
        stats = {"file": path.name, "rows": 0, "missing_added": 0, "excluded_seen": 0}

        for row in rows:
            stats["rows"] += 1
            key = norm(row[key_i])
            en = norm(row[en_i])
            if not key or not en:
                continue
            if key in excluded:
                stats["excluded_seen"] += 1
                continue
            tr = "" if lang_i is None else norm(row[lang_i])
            if not tr and key not in missing_rows:
                needed.add(en)
                missing_rows[key] = {"key": key, "en": en, "source_file": path.name}
                stats["missing_added"] += 1
        source_stats.append(stats)

    TASK_CACHE.parent.mkdir(parents=True, exist_ok=True)
    with open(TASK_CACHE, "w", encoding="utf-8") as f:
        json.dump(
            {"needed": sorted(needed), "missing_rows": list(missing_rows.items()), "source_stats": source_stats},
            f,
            ensure_ascii=False,
        )
    return needed, missing_rows, source_stats


def load_jsonl():
    ok = {}
    blocked = {}
    stats = {"lines": 0, "ok": 0, "blocked": 0, "copy_source": 0, "dify_nb": 0}
    if not JSONL.exists():
        return ok, blocked, stats

    with open(JSONL, encoding="utf-8") as f:
        for line in f:
            stats["lines"] += 1
            try:
                rec = json.loads(line)
            except Exception:
                continue
            en = rec.get("en_full") or ""
            tr = rec.get("translation") or ""
            status = rec.get("status")
            engine = rec.get("engine")
            if engine in stats:
                stats[engine] += 1
            if status == "ok" and en and qc_ok(en, tr):
                ok[en] = tr
                stats["ok"] += 1
            elif en:
                blocked[en] = {"translation": tr, "issues": rec.get("issues") or ["JSONL_BLOCKED"]}
                if status == "blocked":
                    stats["blocked"] += 1
    return ok, blocked, stats


def append_record(record):
    with _buffer_lock:
        _pending_records.append(record)
        should_flush = len(_pending_records) >= FLUSH_EVERY
    if should_flush:
        return flush_records(force=True)
    return 0


def flush_records(force=False):
    with _buffer_lock:
        if not _pending_records:
            return 0
        if not force and len(_pending_records) < FLUSH_EVERY:
            return 0
        JSONL.parent.mkdir(parents=True, exist_ok=True)
        with open(JSONL, "a", encoding="utf-8") as f:
            for record in _pending_records:
                f.write(json.dumps(record, ensure_ascii=False) + "\n")
        count = len(_pending_records)
        _pending_records.clear()
        return count


def check_proxy():
    if not DIFY_API_KEY:
        raise RuntimeError("缺少 DIFY_API_KEY 环境变量，不能调用 Dify 翻译接口")
    r = requests.get(DIFY_HEALTH, timeout=3)
    r.raise_for_status()
    if r.text.strip() != "OK":
        raise RuntimeError(f"Dify proxy health 异常: {r.text[:100]}")


def dify_translate_nb(en):
    safe, mapping = protect(en)
    payload = {
        "inputs": {
            "OriginalLanguage": "English",
            "OriginalText": escape_curly(safe),
            "TargetLanguage": TARGET_CODE,
        },
        "response_mode": "blocking",
        "user": "nb-bokmal-dify-20260617",
    }
    last_error = None
    headers = dict(DIFY_HEADERS_BASE)
    headers["Authorization"] = f"Bearer {DIFY_API_KEY}"
    for attempt in range(3):
        try:
            r = requests.post(
                DIFY_PROXY,
                headers=headers,
                data=json.dumps(payload, ensure_ascii=False),
                timeout=180,
            )
            r.raise_for_status()
            out = r.json().get("data", {}).get("outputs", {}).get("output", [])
            for item in out:
                code = item.get("target_language_code") or ""
                if code.lower() in {"nb", "nb-no"}:
                    tr = unescape_curly(item.get("translated_text", "") or "")
                    return restore(tr, mapping), code, item.get("language_name") or ""
            if out:
                item = out[0]
                tr = unescape_curly(item.get("translated_text", "") or "")
                return restore(tr, mapping), item.get("target_language_code") or "", item.get("language_name") or ""
            return "", "", ""
        except Exception as exc:
            last_error = exc
            time.sleep(min(12, 2 * (attempt + 1) + random.uniform(0.2, 0.8)))
    raise RuntimeError(f"Dify nb failed: {last_error}")


def make_record(en, translation, status, engine, **extra):
    record = {
        "en_hash": md5(en),
        "en_full": en,
        "translation": translation,
        "status": status,
        "engine": engine,
        "target_language": TARGET_CODE,
        "target_name": TARGET_NAME,
        "ts": datetime.now(timezone.utc).isoformat(),
    }
    record.update(extra)
    return record


def translate_one(en):
    if should_copy_source(en):
        rec = make_record(en, en, "ok", "copy_source")
        append_record(rec)
        return rec

    try:
        tr, returned_code, language_name = dify_translate_nb(en)
        issues = qc_issues(en, tr)
        if not issues:
            rec = make_record(
                en,
                tr,
                "ok",
                "dify_nb",
                returned_code=returned_code,
                returned_language_name=language_name,
            )
        else:
            rec = make_record(
                en,
                tr,
                "blocked",
                "dify_nb",
                returned_code=returned_code,
                returned_language_name=language_name,
                issues=issues,
            )
    except Exception as exc:
        rec = make_record(en, "", "blocked", "dify_nb", issues=[f"DIFY_FAILED: {exc}"])
    append_record(rec)
    return rec


def action_probe():
    check_proxy()
    for text in ["Checkout now", "Your order total is {0}.", "<b>Free shipping</b>"]:
        tr, code, name = dify_translate_nb(text)
        print(json.dumps({"en": text, "code": code, "name": name, "translation": tr}, ensure_ascii=False))


def action_stats(force_refresh=False):
    needed, missing_rows, source_stats = build_needed_and_rows(force_refresh=force_refresh)
    ok, blocked, jsonl_stats = load_jsonl()
    pending = [en for en in needed if en not in ok]
    print(json.dumps(
        {
            "source_date": SOURCE_DATE,
            "run_tag": RUN_TAG,
            "jsonl": str(JSONL),
            "task_cache": str(TASK_CACHE),
            "output_xlsx": str(OUT_XLSX),
            "source_stats": source_stats,
            "missing_rows": len(missing_rows),
            "needed_unique_en": len(needed),
            "jsonl_ok": len(ok),
            "jsonl_blocked": len(blocked),
            "pending_unique_en": len(pending),
            "jsonl_stats": jsonl_stats,
        },
        ensure_ascii=False,
        indent=2,
    ))


def action_translate(limit, workers, force_refresh=False):
    check_proxy()
    needed, _missing_rows, source_stats = build_needed_and_rows(force_refresh=force_refresh)
    ok, _blocked, jsonl_stats = load_jsonl()
    pending = [en for en in sorted(needed) if en not in ok]
    if limit:
        pending = pending[:limit]

    log(f"source_stats={json.dumps(source_stats, ensure_ascii=False)}")
    log(f"needed={len(needed)} jsonl_ok={len(ok)} pending={len(pending)} jsonl_stats={jsonl_stats}")
    if not pending:
        return

    counters = {"ok": 0, "blocked": 0, "copy_source": 0, "dify_nb": 0}
    try:
        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(translate_one, en): en for en in pending}
            for idx, future in enumerate(as_completed(futures), 1):
                rec = future.result()
                if rec["status"] == "ok":
                    counters["ok"] += 1
                else:
                    counters["blocked"] += 1
                engine = rec.get("engine")
                if engine in counters:
                    counters[engine] += 1
                if idx <= 3 or idx % 200 == 0:
                    flushed = flush_records(force=False)
                    log(f"progress {idx}/{len(pending)} flushed={flushed} counters={json.dumps(counters, ensure_ascii=False)}")
    finally:
        flushed = flush_records(force=True)
        if flushed:
            log(f"final_flush={flushed}")
    log(f"finished counters={json.dumps(counters, ensure_ascii=False)} jsonl={JSONL}")


def set_text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"
    if cell.value.startswith(FORMULA_PREFIXES):
        cell.quotePrefix = True


def build_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = LANG
    ws.freeze_panes = "A3"
    ws.merge_cells("A1:D1")
    set_text(ws["A1"], NOTICE)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 165

    headers = ["key ID", "en", LANG, "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(2, col)
        set_text(cell, header)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, item in enumerate(rows, 3):
        for col_idx, value in enumerate([item["key"], item["en"], item["translation"], item["remark"]], 1):
            cell = ws.cell(row_idx, col_idx)
            set_text(cell, value)
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if item["issues"]:
            ws.cell(row_idx, 3).fill = RISK_FILL
            ws.cell(row_idx, 4).fill = RISK_FILL

    for col, width in {"A": 28, "B": 62, "C": 62, "D": 46}.items():
        ws.column_dimensions[col].width = width

    UPLOAD.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)


def verify_upload(rows):
    duplicate_keys = []
    seen = set()
    empty_target_unmarked = []
    qc_unmarked = []
    marked_rows = 0
    for item in rows:
        key = item["key"]
        if key in seen:
            duplicate_keys.append(key)
        seen.add(key)
        if item["remark"]:
            marked_rows += 1
        if not item["translation"].strip() and not item["remark"]:
            empty_target_unmarked.append(key)
        row_qc = qc_issues(item["en"], item["translation"])
        if row_qc and not item["remark"]:
            qc_unmarked.append({"key": key, "issues": row_qc})

    issues = []
    if duplicate_keys:
        issues.append(f"DUPLICATE_KEYS count={len(duplicate_keys)} sample={duplicate_keys[:10]}")
    if empty_target_unmarked:
        issues.append(f"EMPTY_TARGET_UNMARKED count={len(empty_target_unmarked)} sample={empty_target_unmarked[:10]}")
    if qc_unmarked:
        issues.append(f"QC_UNMARKED count={len(qc_unmarked)} sample={qc_unmarked[:5]}")
    return {
        "pass": not issues,
        "issues": issues,
        "rows": len(rows),
        "marked_rows": marked_rows,
        "duplicate_key_count": len(duplicate_keys),
        "empty_target_unmarked_count": len(empty_target_unmarked),
        "qc_unmarked_count": len(qc_unmarked),
    }


def action_build_upload(force_refresh=False):
    _needed, missing_rows, source_stats = build_needed_and_rows(force_refresh=force_refresh)
    ok, blocked, jsonl_stats = load_jsonl()
    rows = []
    source_counts = {"jsonl": 0, "blocked_or_missing": 0}

    for item in missing_rows.values():
        en = item["en"]
        translation = ""
        remark = ""
        issues = []
        source = "blocked_or_missing"

        if should_copy_source(en):
            translation = en
            source = "jsonl"
        elif en in ok:
            translation = ok[en]
            source = "jsonl"
        elif en in blocked:
            translation = blocked[en]["translation"]
            issues = blocked[en]["issues"]
        else:
            issues = ["NO_TRANSLATION"]

        if not issues:
            issues = qc_issues(en, translation)
        if issues:
            if len(en) >= 5000 or any(issue.startswith("POSSIBLE_TRUNCATED_LONG_TEXT") for issue in issues):
                remark = "长文本疑似未翻译完整，项目最后统一处理；请勿直接上传该行"
            else:
                remark = f"质检风险：{';'.join(issues)}；请勿直接上传该行"

        source_counts[source] += 1
        rows.append({"key": item["key"], "en": en, "translation": translation, "remark": remark, "issues": issues})

    build_workbook(rows)
    qc = verify_upload(rows)
    report = {
        "lang": LANG,
        "target": TARGET_NAME,
        "output": str(OUT_XLSX),
        "jsonl": str(JSONL),
        "source_files": source_stats,
        "missing_rows": len(missing_rows),
        "source_counts": source_counts,
        "jsonl_stats": jsonl_stats,
        "qc": qc,
        "marked_samples": [
            {
                "key": item["key"],
                "en_len": len(item["en"]),
                "translation_len": len(item["translation"]),
                "remark": item["remark"],
                "issues": item["issues"],
            }
            for item in rows
            if item["issues"]
        ][:50],
    }
    with open(QC_JSON, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--action", choices=["probe", "stats", "translate", "build-upload"], default="probe")
    parser.add_argument("--limit", type=int, default=0, help="translate 小样本限制；0 表示全量")
    parser.add_argument("--workers", type=int, default=3)
    parser.add_argument("--force-refresh", action="store_true", help="重建任务缓存")
    args = parser.parse_args()

    if args.action == "probe":
        action_probe()
    elif args.action == "stats":
        action_stats(force_refresh=args.force_refresh)
    elif args.action == "translate":
        action_translate(limit=args.limit, workers=args.workers, force_refresh=args.force_refresh)
    elif args.action == "build-upload":
        action_build_upload(force_refresh=args.force_refresh)


if __name__ == "__main__":
    main()
